"""Юнит-тесты backend/app/extractors.py: декодирование текста и .docx/.txt."""
from __future__ import annotations

import io

import pytest
from docx import Document

from backend.app import extractors


def _docx_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDecodeText:
    def test_utf8_with_bom(self):
        assert extractors._decode_text("Привет, мир".encode("utf-8-sig")) == "Привет, мир"

    def test_utf16_le_with_bom(self):
        assert (
            extractors._decode_text("Проверка кодировки".encode("utf-16"))
            == "Проверка кодировки"
        )

    def test_cp1251_cyrillic(self):
        data = "Тест 1251".encode("cp1251")
        # нечётное число байтов: utf-16 на таких данных падает и не «перехватывает»
        assert len(data) % 2 == 1
        assert extractors._decode_text(data) == "Тест 1251"

    @pytest.mark.parametrize("data", [b"\x98\xff\x98", b""])
    def test_undecodable_raises(self, data):
        with pytest.raises(extractors.ExtractionError):
            extractors._decode_text(data)


class TestDocx:
    def test_header_body_footer_all_present(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Шапка"
        doc.sections[0].footer.paragraphs[0].text = "Подвал"
        doc.add_paragraph("Тело")
        text = extractors.extract_text("x.docx", _docx_bytes(doc))
        assert "Шапка" in text
        assert "Тело" in text
        assert "Подвал" in text
        assert text.index("Шапка") < text.index("Тело") < text.index("Подвал")

    def test_empty_docx_raises(self):
        with pytest.raises(extractors.ExtractionError, match="не содержит текста"):
            extractors.extract_text("x.docx", _docx_bytes(Document()))

    def test_corrupted_docx_raises(self):
        with pytest.raises(extractors.ExtractionError):
            extractors.extract_text("x.docx", b"\x00\x01not a docx at all")


class TestPlainText:
    @pytest.mark.parametrize("encoding", ["utf-8", "cp1251"])
    def test_txt_encodings(self, encoding):
        assert (
            extractors.extract_text("note.txt", "Русский текст".encode(encoding))
            == "Русский текст"
        )

    def test_md_extension(self):
        assert (
            extractors.extract_text("note.md", "Текст md".encode("utf-8"))
            == "Текст md"
        )


class TestUnsupported:
    def test_unknown_extension_raises(self):
        with pytest.raises(extractors.ExtractionError, match="Неподдерживаемый формат"):
            extractors.extract_text("archive.exe", b"data")

    def test_no_extension_raises(self):
        with pytest.raises(extractors.ExtractionError):
            extractors.extract_text("README", b"data")
