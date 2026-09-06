"""Юнит-тесты чистых функций backend/app/main.py (без HTTP и LLM)."""
from __future__ import annotations

import io
import json

import pytest
from docx import Document

from backend.app import main


def _edits_block(verdict="risk"):
    """JSON-блок с правками: ключ edits первым (важно для brace-matching пути)."""
    block = {
        "edits": [
            {
                "id": "e1",
                "title": "Режим работы",
                "original": "ночная смена",
                "replacement": "дневная смена",
                "reason": "ТК РФ",
            }
        ]
    }
    if verdict is not None:
        block["verdict"] = verdict
    return json.dumps(block, ensure_ascii=False)


class TestParseEdits:
    def test_fenced_block_with_verdict(self):
        report = "## Отчёт\n\nИнструкция проверена.\n\n```json\n" + _edits_block("risk") + "\n```"
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == "## Отчёт\n\nИнструкция проверена."
        assert verdict == "risk"
        assert [e["id"] for e in edits] == ["e1"]
        assert edits[0]["original"] == "ночная смена"

    def test_fenced_block_without_verdict_gives_none(self):
        report = "Отчёт.\n\n```json\n" + _edits_block(verdict=None) + "\n```"
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == "Отчёт."
        assert len(edits) == 1
        assert verdict is None

    def test_invalid_verdict_is_ignored(self):
        report = "Отчёт.\n\n```json\n" + _edits_block(verdict="warning") + "\n```"
        _, edits, verdict = main._parse_edits(report)
        assert len(edits) == 1
        assert verdict is None

    def test_invalid_json_in_fence_report_untouched(self):
        report = "Отчёт.\n\n```json\n{не json}\n```"
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == report
        assert edits == []
        assert verdict is None

    def test_bare_block_orphan_opening_fence_removed(self):
        # открывающий fence есть, закрывающего нет: regex-путь не срабатывает,
        # блок достаётся brace-matching'ом, осиротевший ``` удаляется
        report = "Отчёт.\n```json\n" + _edits_block("risk")
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == "Отчёт."
        assert [e["id"] for e in edits] == ["e1"]
        assert verdict == "risk"
        assert "```" not in cleaned

    def test_bare_block_orphan_closing_fence_removed(self):
        report = _edits_block("risk") + "\n```"
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == ""
        assert [e["id"] for e in edits] == ["e1"]
        assert verdict == "risk"

    def test_bare_block_without_fences_removed_from_middle(self):
        report = "До блока.\n" + _edits_block("fail") + "\nПосле блока."
        cleaned, edits, verdict = main._parse_edits(report)
        assert cleaned == "До блока.\n\nПосле блока."
        assert "edits" not in cleaned
        assert len(edits) == 1
        assert verdict == "fail"

    def test_no_block_returns_report_as_is(self):
        report = "Простой отчёт без structured-блока."
        assert main._parse_edits(report) == (report, [], None)


class TestMatchingBrace:
    def test_nested_braces(self):
        s = '{"a": {"b": {"c": 1}}}'
        assert main._matching_brace(s, 0) == len(s)

    def test_brace_inside_string_ignored(self):
        s = '{"a": "текст с } внутри"}'
        assert main._matching_brace(s, 0) == len(s)

    def test_escaped_quote_inside_string(self):
        s = '{"a": "x \\" y", "b": 2}'
        assert main._matching_brace(s, 0) == len(s)

    def test_escaped_brace_inside_string(self):
        # \{ внутри строки не открывает вложенность; закрывающая скобка — перед "y"
        s = '{"a": "x\\{"}y"}'
        assert main._matching_brace(s, 0) == s.index("}y") + 1

    @pytest.mark.parametrize("s", ['{"a": 1', "{", '{"a": "незакрытая строка}}'])
    def test_unclosed_returns_minus_one(self, s):
        assert main._matching_brace(s, 0) == -1


class TestNormalizeEdits:
    def test_skips_non_dict_and_generates_ids(self):
        data = {
            "edits": [
                "мусор",
                {"title": "Без id"},
                None,
                {"id": "e9", "original": "цитата"},
            ]
        }
        assert main._normalize_edits(data) == [
            {"id": "e2", "title": "Без id", "original": "", "replacement": "", "reason": ""},
            {"id": "e9", "title": "", "original": "цитата", "replacement": "", "reason": ""},
        ]

    def test_values_coerced_to_str(self):
        data = {
            "edits": [
                {"id": 7, "title": None, "original": 1.5, "replacement": True, "reason": ["а"]}
            ]
        }
        (edit,) = main._normalize_edits(data)
        assert edit == {
            "id": "7",
            "title": "",
            "original": "1.5",
            "replacement": "True",
            "reason": "['а']",
        }

    @pytest.mark.parametrize(
        "data",
        [None, [{"edits": []}], {"verdict": "ok"}, {"edits": "не список"}],
    )
    def test_invalid_structure_returns_none(self, data):
        assert main._normalize_edits(data) is None


class TestSummaryVerdict:
    def test_fail_phrase_wins(self):
        assert main._summary_verdict("Документ НЕ СООТВЕТСТВУЕТ требованиям") == "fail"
        assert main._summary_verdict("риск: инструкция не соответствует") == "fail"

    def test_risk(self):
        assert main._summary_verdict("Выявлен РИСК нарушения") == "risk"

    def test_ok(self):
        assert main._summary_verdict("Полностью соответствует") == "ok"


class TestMarkdownToDocx:
    def _doc(self, text):
        return Document(io.BytesIO(main._markdown_to_docx(text)))

    def test_headings_and_lists(self):
        doc = self._doc(
            "# Заголовок 1\n## Заголовок 2\n### Заголовок 3\n"
            "- пункт маркированный\n1. пункт нумерованный\nОбычный абзац"
        )
        styles = [(p.style.name, p.text) for p in doc.paragraphs]
        assert styles == [
            ("Heading 1", "Заголовок 1"),
            ("Heading 2", "Заголовок 2"),
            ("Heading 3", "Заголовок 3"),
            ("List Bullet", "пункт маркированный"),
            ("List Number", "пункт нумерованный"),
            ("Normal", "Обычный абзац"),
        ]

    def test_inline_markdown_cleaned(self):
        doc = self._doc("**Жирный** текст и `код`")
        assert [p.text for p in doc.paragraphs] == ["Жирный текст и код"]

    @pytest.mark.parametrize("fence", ["```markdown", "```"])
    def test_code_fence_unwrapped(self, fence):
        doc = self._doc(fence + "\n# Заголовок\nТекст\n```")
        assert [(p.style.name, p.text) for p in doc.paragraphs] == [
            ("Heading 1", "Заголовок"),
            ("Normal", "Текст"),
        ]

    def test_into_existing_doc_appends(self):
        doc = Document()
        doc.add_paragraph("Вступление")
        main._markdown_into_docx(doc, "# Раздел\n\n- пункт\n\nТекст раздела")
        styles = [(p.style.name, p.text) for p in doc.paragraphs]
        assert styles[0] == ("Normal", "Вступление")
        assert ("Heading 1", "Раздел") in styles
        assert ("List Bullet", "пункт") in styles
        assert ("Normal", "Текст раздела") in styles


def _new_paragraph(doc, *run_texts):
    par = doc.add_paragraph()
    for text in run_texts:
        par.add_run(text)
    return par


class TestReplaceInParagraph:
    def test_single_run_replaced(self):
        doc = Document()
        par = _new_paragraph(doc, "Сторож обязан дежурить ночью ежедневно")
        assert main._replace_in_paragraph(par, "дежурить ночью", "работать днём") is True
        assert par.text == "Сторож обязан работать днём ежедневно"

    def test_replacement_across_runs(self):
        doc = Document()
        par = _new_paragraph(doc, "начало ", "ЛОМАН", "НАЯ конец")
        assert main._replace_in_paragraph(par, "ЛОМАННАЯ", "X") is True
        assert par.text == "начало X конец"

    def test_text_outside_range_preserved(self):
        doc = Document()
        par = _new_paragraph(doc, "ДО ", "ЦЕЛЬ", " ПОСЛЕ")
        assert main._replace_in_paragraph(par, "ЦЕЛЬ", "замена") is True
        assert [r.text for r in par.runs] == ["ДО ", "замена", " ПОСЛЕ"]

    def test_not_found_returns_false(self):
        doc = Document()
        par = _new_paragraph(doc, "Просто текст")
        assert main._replace_in_paragraph(par, "нет такой", "замена") is False
        assert par.text == "Просто текст"


class TestApplyEditsToDocx:
    EDIT = {
        "id": "e1",
        "title": "t",
        "original": "ночное патрулирование",
        "replacement": "НОВАЯ ФОРМУЛИРОВКА",
        "reason": "r",
    }

    def _docx(self, doc):
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_body_quote_applied(self):
        doc = Document()
        doc.add_paragraph("Сторож выполняет ночное патрулирование территории.")
        result, failed = main._apply_edits_to_docx(self._docx(doc), [dict(self.EDIT)])
        assert failed == []
        assert result is not None
        text = "\n".join(p.text for p in result.paragraphs)
        assert "НОВАЯ ФОРМУЛИРОВКА" in text
        assert "ночное патрулирование" not in text

    def test_table_quote_applied(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Особые обязанности: ночное патрулирование"
        result, failed = main._apply_edits_to_docx(self._docx(doc), [dict(self.EDIT)])
        assert failed == []
        cells = "\n".join(
            cell.text for row in result.tables[0].rows for cell in row.cells
        )
        assert "НОВАЯ ФОРМУЛИРОВКА" in cells

    def test_header_quote_applied(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "УТВЕРЖДАЮ: ночное патрулирование"
        doc.add_paragraph("Тело инструкции")
        result, failed = main._apply_edits_to_docx(self._docx(doc), [dict(self.EDIT)])
        assert failed == []
        header_text = "\n".join(
            p.text for p in result.sections[0].header.paragraphs
        )
        assert "НОВАЯ ФОРМУЛИРОВКА" in header_text

    def test_missing_quote_reported_as_failed(self):
        doc = Document()
        doc.add_paragraph("Совсем другой текст.")
        edit = dict(self.EDIT, original="такой цитаты точно нет 12345")
        result, failed = main._apply_edits_to_docx(self._docx(doc), [edit])
        assert failed == [edit]
        assert result is not None


class TestBuildUserMessage:
    def test_all_fields(self):
        msg = main._build_user_message(
            "di.docx", "ТЕКСТ ДИ", "охрана объектов", "full", "ночные смены"
        )
        assert msg.startswith("Должностная инструкция (файл: di.docx):\n\nТЕКСТ ДИ")
        assert "Предмет госконтракта: охрана объектов" in msg
        assert (
            "Занятость по контракту: сотрудник работает полностью по контракту" in msg
        )
        assert (
            "Дополнительный контекст (требования к персоналу и т.п.): ночные смены"
            in msg
        )
        assert msg.index("ТЕКСТ ДИ") < msg.index("Предмет госконтракта")

    def test_minimal_only_required(self):
        msg = main._build_user_message("f.txt", "Тело", "", "", "")
        assert msg == "Должностная инструкция (файл: f.txt):\n\nТело"
        assert "---" not in msg

    @pytest.mark.parametrize(
        ("employment_type", "present", "absent"),
        [
            ("full", "полностью по контракту", "частично"),
            ("partial", "по контракту частично", "полностью"),
        ],
    )
    def test_employment_type_variants(self, employment_type, present, absent):
        msg = main._build_user_message("f", "t", "", employment_type, "")
        assert present in msg
        assert absent not in msg

    def test_unknown_employment_type_omitted(self):
        msg = main._build_user_message("f", "t", "", "другое", "")
        assert "Занятость по контракту" not in msg
