"""Тесты «Пересечений обязанностей» (Ф1) и шаблонного экспорта (Ф2).

Всё офлайн: LLM мокается на уровне main.llm.chat_completion, БД и шаблон —
в tmp_path (изоляция через conftest.isolated_settings)."""
from __future__ import annotations

import io
import json
import sqlite3
import zipfile
from datetime import datetime

import pytest
from docx import Document

from backend.app import history, main

_DOC_COUNTER = {"n": 0}


def add_document(text: str, position: str = "", department: str = "") -> int:
    """Документ с check-версией через сервисный слой; возвращает id документа."""
    _DOC_COUNTER["n"] += 1
    batch = history.ensure_batch(
        f"jobovl{_DOC_COUNTER['n']:03d}", "deepseek", "m", "", "", ""
    )
    vid = history.save_check_version(
        batch, "di.docx", ".docx", b"", text, False, "R", "ok", []
    )
    doc_id = history.get_version(vid)["document_id"]
    if position or department:
        history.update_document(doc_id, position=position, department=department)
    return doc_id


def set_duties_current(doc_id: int, duties: list[str]) -> None:
    """Выставляет обязанности вручную на id актуальной check-версии (кэш)."""
    vid = history.latest_check_version(doc_id)[0]
    history.set_duties(doc_id, duties, vid)


def install_llm(monkeypatch, duties_by_marker: dict, overlap_raw: str, calls: list):
    """Мок LLM: извлечение — по маркеру текста, сравнение — по системному промту."""

    async def fake(provider, model, system_prompt, user_message):
        calls.append({"system": system_prompt, "user": user_message})
        if system_prompt == main.DUTIES_SYSTEM_PROMPT:
            for marker, duties in duties_by_marker.items():
                if marker in user_message:
                    return json.dumps({"duties": duties}, ensure_ascii=False)
            return json.dumps({"duties": ["Общая обязанность"]}, ensure_ascii=False)
        assert system_prompt == main.OVERLAPS_SYSTEM_PROMPT
        return overlap_raw

    monkeypatch.setattr(main.llm, "chat_completion", fake)


def docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


def build_template_docx() -> bytes:
    """Шаблон со всеми ключами TEMPLATE_KEYS."""
    doc = Document()
    doc.add_paragraph("Приложение к приказу")
    doc.add_paragraph("Должность: {{ДОЛЖНОСТЬ}}")
    doc.add_paragraph("Подразделение: {{ПОДРАЗДЕЛЕНИЕ}}")
    doc.add_paragraph("Наименование: {{НАЗВАНИЕ}}")
    doc.add_paragraph("Дата: {{ДАТА}}")
    doc.add_paragraph("Заметки: {{ЗАМЕТКИ}}")
    doc.add_paragraph("{{ТЕКСТ}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Миграция схемы ----------

def test_migration_adds_duty_columns_and_keeps_data():
    db = history._db_path()
    db.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db) as conn:
        # старая схема (до фичи «пересечения»): documents без duties-колонок
        conn.executescript(
            """
            CREATE TABLE documents (
                id INTEGER PRIMARY KEY,
                position TEXT NOT NULL DEFAULT '',
                department TEXT NOT NULL DEFAULT '',
                title TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            INSERT INTO documents(title, created_at, updated_at)
            VALUES ('старый документ', '2026-01-01T00:00:00.000+00:00',
                    '2026-01-01T00:00:00.000+00:00');
            """
        )
        before = {r[1] for r in conn.execute("PRAGMA table_info(documents)")}
    assert "duties_json" not in before

    total, items = history.list_documents("", 10, 0)  # триггерит инициализацию схемы
    assert total == 1
    assert items[0]["title"] == "старый документ"

    with sqlite3.connect(db) as conn:
        after = {r[1] for r in conn.execute("PRAGMA table_info(documents)")}
        tables = {
            r[0]
            for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
        }
    assert {"duties_json", "duties_version_id", "duties_extracted_at"} <= after
    assert "overlap_reports" in tables

    # данные живы, новые колонки работают
    doc_id = items[0]["id"]
    history.set_duties(doc_id, ["обязанность из новой версии"], 7)
    duties, vid, extracted_at = history.get_duties(doc_id)
    assert duties == ["обязанность из новой версии"]
    assert vid == 7
    assert extracted_at
    total, items = history.list_documents("", 10, 0)
    assert total == 1
    assert items[0]["title"] == "старый документ"


def test_set_get_duties_roundtrip():
    doc_id = add_document("Текст обязанностей.", "Сторож", "Охрана")
    assert history.get_duties(doc_id) == (None, None, None)
    history.set_duties(doc_id, [], 1)  # пустой список не пишет и не затирает
    assert history.get_duties(doc_id) == (None, None, None)
    history.set_duties(doc_id, ["дежурство", "осмотр"], 3)
    duties, vid, ts = history.get_duties(doc_id)
    assert duties == ["дежурство", "осмотр"]
    assert vid == 3
    assert ts
    history.set_duties(doc_id, ["новая обязанность"], 5)
    duties, vid, _ = history.get_duties(doc_id)
    assert duties == ["новая обязанность"]
    assert vid == 5


# ---------- extract-duties ----------

def test_extract_duties_success(client, monkeypatch):
    doc_id = add_document(
        "Должностная инструкция сторожа. Обязанности: ...", "Сторож", "Охрана"
    )
    calls = []

    async def fake(provider, model, system_prompt, user_message):
        calls.append({"system": system_prompt, "user": user_message})
        assert system_prompt == main.DUTIES_SYSTEM_PROMPT
        assert user_message.startswith("Должностная инструкция сторожа")
        # мусор вокруг JSON — должен отработать brace-matching
        return (
            "Конечно!\n```json\n"
            '{"duties": ["дежурство на посту", "осмотр территории"]}\n```'
        )

    monkeypatch.setattr(main.llm, "chat_completion", fake)
    resp = client.post(f"/api/history/documents/{doc_id}/extract-duties")
    assert resp.status_code == 200
    vid = history.latest_check_version(doc_id)[0]
    assert resp.json() == {
        "documentId": doc_id,
        "count": 2,
        "duties": ["дежурство на посту", "осмотр территории"],
    }
    duties, saved_vid, extracted_at = history.get_duties(doc_id)
    assert duties == ["дежурство на посту", "осмотр территории"]
    assert saved_vid == vid
    assert extracted_at


def test_extract_duties_invalid_response_502(client, monkeypatch):
    doc_id = add_document("Текст для невалидного ответа.", "Сторож", "Охрана")

    async def fake(provider, model, system_prompt, user_message):
        return "ответ вообще без JSON"

    monkeypatch.setattr(main.llm, "chat_completion", fake)
    resp = client.post(f"/api/history/documents/{doc_id}/extract-duties")
    assert resp.status_code == 502
    assert resp.json()["detail"] == "invalid_duties_response"
    assert history.get_duties(doc_id) == (None, None, None)


def test_extract_duties_no_text_and_missing_doc(client, monkeypatch):
    empty_doc = add_document("", "Сторож", "Охрана")

    async def fail(provider, model, system_prompt, user_message):
        raise AssertionError("LLM не должен вызываться")

    monkeypatch.setattr(main.llm, "chat_completion", fail)

    resp = client.post(f"/api/history/documents/{empty_doc}/extract-duties")
    assert resp.status_code == 409
    assert resp.json()["detail"] == "no_text"

    missing = client.post("/api/history/documents/99999/extract-duties")
    assert missing.status_code == 404
    assert missing.json()["detail"] == "document_not_found"


def test_parse_duties_caps():
    raw = json.dumps(
        {"duties": [f"обязанность {i}" for i in range(50)] + ["", "   "]},
        ensure_ascii=False,
    )
    duties = main._parse_duties_response(raw)
    assert len(duties) == main.DUTIES_MAX_COUNT
    long = json.dumps({"duties": ["х" * 500]}, ensure_ascii=False)
    assert len(main._parse_duties_response(long)[0]) == main.DUTIES_MAX_LEN
    for bad in ('{"duties": []}', '{"duties": [42]}', "не json", '{"other": 1}'):
        with pytest.raises(main.llm.LLMError):
            main._parse_duties_response(bad)


# ---------- overlaps ----------

def test_overlaps_end_to_end_and_get(client, monkeypatch):
    doc_a = add_document("Текст А: контроль проходной.", "Сторож", "Служба охраны")
    doc_b = add_document("Текст Б: охрана объекта.", "Инженер", "Служба эксплуатации")
    doc_c = add_document("Текст В: снабжение.", "Завхоз", "Служба эксплуатации")

    overlap_raw = json.dumps(
        {
            "groups": [
                {
                    "duty": "Проверка пропускного режима",
                    "comment": "дублирование зоны ответственности",
                    "items": [
                        {"documentId": doc_a, "duty": "контролировать проходную"},
                        {"documentId": doc_b, "duty": "обеспечивать охрану"},
                    ],
                },
                # одиночный item — выбрасывается
                {"duty": "одиночное", "comment": "",
                 "items": [{"documentId": doc_a, "duty": "x"}]},
                # одно подразделение — выбрасывается
                {"duty": "внутри отдела", "comment": "",
                 "items": [{"documentId": doc_b, "duty": "x"},
                           {"documentId": doc_c, "duty": "y"}]},
                # чужой documentId → остаётся 1 item — выбрасывается
                {"duty": "чужой id", "comment": "",
                 "items": [{"documentId": 99999, "duty": "x"},
                           {"documentId": doc_a, "duty": "y"}]},
            ]
        },
        ensure_ascii=False,
    )
    calls = []
    install_llm(
        monkeypatch,
        {
            "Текст А": ["контролировать проходную", "вести журнал"],
            "Текст Б": ["обеспечивать охрану", "ремонт оборудования"],
            "Текст В": ["обеспечивать снабжение", "учёт имущества"],
        },
        overlap_raw,
        calls,
    )

    # до первого отчёта — 404 no_report
    empty = client.get("/api/history/overlaps")
    assert empty.status_code == 404
    assert empty.json()["detail"] == "no_report"

    resp = client.post("/api/history/overlaps", json={"documentIds": None})
    assert resp.status_code == 200
    data = resp.json()
    assert isinstance(data["id"], int)
    assert data["createdAt"]
    assert {d["documentId"] for d in data["documents"]} == {doc_a, doc_b, doc_c}
    for d in data["documents"]:
        assert d["dutiesCount"] == 2
        assert d["position"] and d["department"]
    assert data["skipped"] == []
    assert len(data["groups"]) == 1
    group = data["groups"][0]
    assert group["duty"] == "Проверка пропускного режима"
    assert group["comment"] == "дублирование зоны ответственности"
    assert {(i["documentId"], i["position"], i["department"]) for i in group["items"]} == {
        (doc_a, "Сторож", "Служба охраны"),
        (doc_b, "Инженер", "Служба эксплуатации"),
    }

    # GET возвращает тот же (последний) отчёт
    assert client.get("/api/history/overlaps").json() == data

    # извлечение было по одному вызову на документ + один вызов сравнения
    assert len([c for c in calls if c["system"] == main.DUTIES_SYSTEM_PROMPT]) == 3
    assert len([c for c in calls if c["system"] == main.OVERLAPS_SYSTEM_PROMPT]) == 1


def test_overlaps_skipped_entries(client, monkeypatch):
    ok_doc = add_document("Текст успешного документа.", "Сторож", "Охрана")
    empty_doc = add_document("", "Сторож без текста", "Охрана")
    calls = []
    install_llm(
        monkeypatch,
        {"Текст успешного": ["дежурство"]},
        json.dumps({"groups": []}, ensure_ascii=False),
        calls,
    )
    resp = client.post(
        "/api/history/overlaps", json={"documentIds": [ok_doc, empty_doc, 99999]}
    )
    assert resp.status_code == 200
    data = resp.json()
    assert [d["documentId"] for d in data["documents"]] == [ok_doc]
    by_id = {s["documentId"]: s for s in data["skipped"]}
    assert by_id[empty_doc]["error"] == "no_text"
    assert by_id[99999]["error"] == "document_not_found"
    assert by_id[99999]["position"] == ""
    assert by_id[99999]["department"] == ""
    assert data["groups"] == []


def test_overlaps_extraction_failure_skipped(client, monkeypatch):
    doc_id = add_document("Текст с ошибкой извлечения.", "Сторож", "Охрана")

    async def fail(provider, model, system_prompt, user_message):
        raise main.llm.LLMError("модель недоступна")

    monkeypatch.setattr(main.llm, "chat_completion", fail)
    resp = client.post("/api/history/overlaps", json={"documentIds": [doc_id]})
    assert resp.status_code == 200
    data = resp.json()
    assert data["documents"] == []
    assert data["groups"] == []
    assert data["skipped"][0]["error"] == "извлечение не удалось: модель недоступна"


def test_overlaps_empty_set_400(client):
    resp = client.post("/api/history/overlaps", json={"documentIds": None})
    assert resp.status_code == 400
    assert resp.json()["detail"] == "no_documents"
    resp = client.post("/api/history/overlaps", json={"documentIds": []})
    assert resp.status_code == 400
    assert resp.json()["detail"] == "no_documents"


def test_overlaps_cache_no_extraction(client, monkeypatch):
    doc_id = add_document("Текст кэш-теста.", "Кладовщик", "Логистика")
    set_duties_current(doc_id, ["приём товара", "учёт тары"])

    calls = []
    install_llm(monkeypatch, {}, json.dumps({"groups": []}, ensure_ascii=False), calls)
    resp = client.post("/api/history/overlaps", json={"documentIds": [doc_id]})
    assert resp.status_code == 200
    assert resp.json()["documents"][0]["dutiesCount"] == 2
    # обязанности взяты из кэша: только один вызов сравнения, извлечения не было
    assert [c["system"] for c in calls] == [main.OVERLAPS_SYSTEM_PROMPT]


def test_overlaps_second_run_uses_cache(client, monkeypatch):
    doc_id = add_document("Текст повторного прогона.", "Сторож", "Охрана")
    calls = []
    install_llm(
        monkeypatch,
        {"Текст повторного": ["дежурство на посту"]},
        json.dumps({"groups": []}, ensure_ascii=False),
        calls,
    )
    first = client.post("/api/history/overlaps", json={"documentIds": [doc_id]})
    assert first.status_code == 200
    assert len([c for c in calls if c["system"] == main.DUTIES_SYSTEM_PROMPT]) == 1

    before = len(calls)
    second = client.post("/api/history/overlaps", json={"documentIds": [doc_id]})
    assert second.status_code == 200
    assert second.json()["id"] == first.json()["id"] + 1
    new_calls = calls[before:]
    # повторный прогон не перезапрашивает извлечение — только сравнение
    assert [c["system"] for c in new_calls] == [main.OVERLAPS_SYSTEM_PROMPT]


def test_overlaps_chunks_of_eight(client, monkeypatch):
    ids = []
    for i in range(9):
        doc_id = add_document(f"Текст номер {i}.", f"Должность {i}", f"Отдел {i}")
        set_duties_current(doc_id, [f"обязанность {i}"])
        ids.append(doc_id)
    calls = []
    install_llm(monkeypatch, {}, json.dumps({"groups": []}, ensure_ascii=False), calls)

    resp = client.post("/api/history/overlaps", json={"documentIds": None})
    assert resp.status_code == 200
    compare = [c for c in calls if c["system"] == main.OVERLAPS_SYSTEM_PROMPT]
    assert len(compare) == 2  # 8 + 1
    assert f"### Документ {ids[0]} —" in compare[0]["user"]
    assert f"### Документ {ids[7]} —" in compare[0]["user"]
    assert f"### Документ {ids[8]} —" not in compare[0]["user"]
    assert f"### Документ {ids[8]} —" in compare[1]["user"]
    # порядок документов в отчёте сохранён
    assert [d["documentId"] for d in resp.json()["documents"]] == ids


def test_overlap_reports_keep_last_five():
    for i in range(1, 7):
        history.save_overlap_report(
            [i], {"documents": [], "groups": [], "skipped": [], "n": i}
        )
    rec = history.latest_overlap_report()
    assert rec is not None
    assert rec["document_ids"] == [6]
    assert rec["result"]["n"] == 6
    with sqlite3.connect(history._db_path()) as conn:
        count = conn.execute("SELECT COUNT(*) FROM overlap_reports").fetchone()[0]
    assert count == history.OVERLAP_REPORTS_LIMIT


def test_overlap_latest_none_initially():
    assert history.latest_overlap_report() is None


# ---------- шаблон ----------

def test_template_put_get_delete(client):
    data = build_template_docx()
    put = client.put(
        "/api/history/template",
        content=data,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert put.status_code == 200
    body = put.json()
    assert body["ok"] is True
    assert body["size"] == len(data)
    assert body["keys"] == [
        "ДОЛЖНОСТЬ", "ПОДРАЗДЕЛЕНИЕ", "НАЗВАНИЕ", "ТЕКСТ", "ДАТА", "ЗАМЕТКИ",
    ]
    assert history.template_path().is_file()

    got = client.get("/api/history/template")
    assert got.status_code == 200
    assert got.json() == {"exists": True, "size": len(data), "keys": body["keys"]}

    assert client.delete("/api/history/template").json() == {"ok": True}
    assert client.get("/api/history/template").json() == {"exists": False}
    # повторное удаление идемпотентно
    assert client.delete("/api/history/template").json() == {"ok": True}


def test_template_put_without_keys(client, make_docx_bytes):
    data = make_docx_bytes(["Обычный документ без плейсхолдеров"])
    put = client.put(
        "/api/history/template",
        content=data,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert put.status_code == 200
    assert put.json() == {"ok": True, "size": len(data), "keys": []}


def test_template_put_invalid(client):
    resp = client.put(
        "/api/history/template",
        content="это точно не docx".encode("utf-8"),
        headers={"Content-Type": "application/octet-stream"},
    )
    assert resp.status_code == 400
    assert resp.json()["detail"] == "invalid_template"
    # zip-магия есть, но python-docx файл не открывает
    fake = b"PK\x03\x04" + "мусор вместо архива".encode("utf-8")
    resp = client.put(
        "/api/history/template",
        content=fake,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert resp.status_code == 400
    assert resp.json()["detail"] == "invalid_template"
    assert history.get_template() is None


def test_assemble_from_template_text_fallback():
    data = main.assemble_from_template(
        build_template_docx(),
        position="Сторож",
        department="Служба охраны",
        title="ДИ сторожа",
        notes="примечание",
        body_docx_bytes=None,
        text="Строка первая\nСтрока вторая",
    )
    text = docx_text(data)
    assert "Должность: Сторож" in text
    assert "Подразделение: Служба охраны" in text
    assert "Наименование: ДИ сторожа" in text
    assert "Заметки: примечание" in text
    assert datetime.now().strftime("%d.%m.%Y") in text
    assert "Строка первая" in text
    assert "Строка вторая" in text
    assert "{{" not in text
    assert "Приложение к приказу" in text  # абзацы без плейсхолдеров не тронуты


def test_assemble_from_template_table_and_empty_values():
    """Плейсхолдеры в ячейках таблицы заменяются; пустые значения → пустая строка."""
    doc = Document()
    doc.add_paragraph("Должность: {{ДОЛЖНОСТЬ}}, дата {{ДАТА}}")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Отдел: {{ПОДРАЗДЕЛЕНИЕ}}"
    table.rows[0].cells[1].text = "Заметки: {{ЗАМЕТКИ}}"
    buf = io.BytesIO()
    doc.save(buf)
    data = main.assemble_from_template(
        buf.getvalue(), position="Сторож", department="Охрана"
    )
    result = Document(io.BytesIO(data))
    cell_texts = [
        c.text for row in result.tables[0].rows for c in row.cells
    ]
    assert cell_texts == ["Отдел: Охрана", "Заметки: "]
    assert "Должность: Сторож, дата " + datetime.now().strftime("%d.%m.%Y") in (
        result.paragraphs[0].text
    )


# ---------- export/fixed?template=true ----------

def _make_doc_with_fix(make_docx_bytes, job_no: str, position: str, department: str):
    batch = history.ensure_batch(f"jobexp{job_no}", "deepseek", "m", "", "", "")
    check_vid = history.save_check_version(
        batch, "di.docx", ".docx", b"", "Исходный текст", False, "R", "risk", []
    )
    doc_id = history.get_version(check_vid)["document_id"]
    history.update_document(doc_id, position=position, department=department)
    return doc_id, check_vid


def test_export_fixed_with_template(client, make_docx_bytes):
    doc_id, check_vid = _make_doc_with_fix(
        make_docx_bytes, "01", "Сторож", "Служба охраны"
    )

    client.put(
        "/api/history/template",
        content=build_template_docx(),
        headers={"Content-Type": "application/octet-stream"},
    )
    # шаблон есть, fix-версий нет → 404
    resp = client.get("/api/history/export/fixed", params={"template": "true"})
    assert resp.status_code == 404
    assert resp.json()["detail"] == "no_fixed_files"

    history.save_fixed_version(
        check_vid,
        make_docx_bytes(["Исправленный абзац с МАРКЕРОМ-ФИКСА."]),
        "docx",
        ["e1"],
        "di_исправленная.docx",
    )

    resp = client.get("/api/history/export/fixed", params={"template": "true"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/zip")
    assert "attachment" in resp.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        assert zf.namelist() == ["Сторож_исправленная.docx"]
        assembled = zf.read("Сторож_исправленная.docx")
    text = docx_text(assembled)
    assert "МАРКЕРОМ-ФИКСА" in text       # содержимое исправленного файла вставлено
    assert "Должность: Сторож" in text    # {{ДОЛЖНОСТЬ}} заменена
    assert "Подразделение: Служба охраны" in text
    assert datetime.now().strftime("%d.%m.%Y") in text
    assert "{{" not in text


def test_export_fixed_template_missing_409(client, make_docx_bytes):
    _, check_vid = _make_doc_with_fix(make_docx_bytes, "02", "Сторож", "Охрана")
    history.save_fixed_version(
        check_vid, make_docx_bytes(["Фикс"]), "docx", [], "di.docx"
    )
    resp = client.get("/api/history/export/fixed", params={"template": "true"})
    assert resp.status_code == 409
    assert resp.json()["detail"] == "no_template"


def test_export_fixed_template_collision_names(client, make_docx_bytes):
    _, vid1 = _make_doc_with_fix(make_docx_bytes, "03", "Сторож", "Охрана-1")
    history.save_fixed_version(
        vid1, make_docx_bytes(["Фикс один"]), "docx", [], "a.docx"
    )
    _, vid2 = _make_doc_with_fix(make_docx_bytes, "04", "Сторож", "Охрана-2")
    history.save_fixed_version(
        vid2, make_docx_bytes(["Фикс два"]), "docx", [], "b.docx"
    )
    client.put(
        "/api/history/template",
        content=build_template_docx(),
        headers={"Content-Type": "application/octet-stream"},
    )
    resp = client.get("/api/history/export/fixed", params={"template": "true"})
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        assert sorted(zf.namelist()) == [
            "Сторож_исправленная.docx",
            "Сторож_исправленная_2.docx",
        ]
        texts = " ".join(docx_text(zf.read(n)) for n in zf.namelist())
    assert "Фикс один" in texts
    assert "Фикс два" in texts


def test_export_fixed_template_by_ids(client, make_docx_bytes):
    doc1, vid1 = _make_doc_with_fix(make_docx_bytes, "05", "Первый", "Отдел 1")
    history.save_fixed_version(vid1, make_docx_bytes(["Фикс первый"]), "docx", [], "a.docx")
    doc2, vid2 = _make_doc_with_fix(make_docx_bytes, "06", "Второй", "Отдел 2")
    history.save_fixed_version(vid2, make_docx_bytes(["Фикс второй"]), "docx", [], "b.docx")

    client.put(
        "/api/history/template",
        content=build_template_docx(),
        headers={"Content-Type": "application/octet-stream"},
    )
    resp = client.get(
        "/api/history/export/fixed",
        params={"template": "true", "documentIds": str(doc2)},
    )
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        assert zf.namelist() == ["Второй_исправленная.docx"]
        assert "Фикс второй" in docx_text(zf.read("Второй_исправленная.docx"))


def test_export_fixed_without_template_previous_behavior(client, make_docx_bytes):
    """Без template=true — прежний zip с нетронутыми fix-байтами."""
    _, check_vid = _make_doc_with_fix(make_docx_bytes, "07", "Сторож", "Охрана")
    history.save_fixed_version(
        check_vid, b"FIXED-BYTES", "docx", ["e1"], "di_исправленная.docx"
    )
    resp = client.get("/api/history/export/fixed")
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        assert zf.namelist() == ["di_исправленная.docx"]
        assert zf.read("di_исправленная.docx") == b"FIXED-BYTES"
