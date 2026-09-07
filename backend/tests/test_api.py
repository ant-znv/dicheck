"""Интеграционные тесты API DI_Check (TestClient, LLM замокан, сети нет)."""
from __future__ import annotations

import asyncio
import io
import json
import time

import pytest
from docx import Document

from backend.app import main, settings


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
QUOTE = "дежурить в ночное время суток"
REPLACEMENT = "НОВАЯ ФОРМУЛИРОВКА"
MISSING_QUOTE = "такой цитаты точно нет в документе 12345"
BODY_PARAGRAPHS = [
    "Должностная инструкция сторожа",
    "1. Общие положения",
    "Сторож подчиняется напрямую директору.",
    "Сторож обязан дежурить в ночное время суток.",
]
TERMINAL_JOB_STATUSES = {"done", "error", "cancelled"}


def structured_response(quote=QUOTE, replacement=REPLACEMENT, verdict="risk"):
    """Ответ «модели»: markdown-отчёт + fenced-блок с verdict и edits."""
    block = {
        "verdict": verdict,
        "edits": [
            {
                "id": "e1",
                "title": "Режим работы",
                "original": quote,
                "replacement": replacement,
                "reason": "Требования трудового договора",
            }
        ],
    }
    return (
        "## Отчёт о проверке\n\n"
        "Инструкция проверена. Найдены замечания.\n\n"
        "```json\n" + json.dumps(block, ensure_ascii=False) + "\n```"
    )


def post_check(client, docx_bytes, filename="di.docx", **form):
    data = {"contractSubject": "", "employmentType": "", "extraContext": ""}
    data.update(form)
    return client.post(
        "/api/check",
        files={"files": (filename, docx_bytes, DOCX_MIME)},
        data=data,
    )


def poll_job(client, job_id, want="done", timeout=5.0):
    """Опрашивает джобу до статуса want; падает на других терминальных статусах."""
    deadline = time.monotonic() + timeout
    last = None
    while time.monotonic() < deadline:
        resp = client.get(f"/api/jobs/{job_id}")
        assert resp.status_code == 200, resp.text
        last = resp.json()
        if last["status"] == want:
            return last
        assert last["status"] not in TERMINAL_JOB_STATUSES, (
            f"джоба завершилась со статусом {last['status']!r}, "
            f"ожидали {want!r}: {json.dumps(last, ensure_ascii=False)}"
        )
        time.sleep(0.05)
    pytest.fail(
        f"таймаут ожидания статуса {want!r}, "
        f"последний статус {last and last['status']!r}"
    )


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


# ---------- 1. Нет API-ключа ----------

def test_check_without_api_key_returns_409(client, monkeypatch, make_docx_bytes):
    monkeypatch.setattr(settings, "has_api_key", lambda provider: False)
    resp = post_check(client, make_docx_bytes(BODY_PARAGRAPHS))
    assert resp.status_code == 409
    assert resp.json()["detail"] == "no_api_key"
    assert main._jobs == {}


# ---------- 2. Полный цикл: check → poll → export ----------

def test_full_cycle_check_poll_export(
    client, with_api_key, make_fake_llm, make_docx_bytes
):
    calls = make_fake_llm([structured_response(verdict="risk")])
    resp = post_check(
        client,
        make_docx_bytes(BODY_PARAGRAPHS),
        contractSubject="охрана объектов",
        employmentType="full",
    )
    assert resp.status_code == 200
    job_id = resp.json()["jobId"]

    data = poll_job(client, job_id, "done")
    result = data["results"][0]
    assert result["status"] == "done"
    assert result["error"] is None
    assert result["summaryVerdict"] == "risk"
    assert result["textTruncated"] is False
    assert [e["id"] for e in result["edits"]] == ["e1"]
    assert result["edits"][0]["original"] == QUOTE
    assert result["edits"][0]["replacement"] == REPLACEMENT
    # отчёт очищен от structured-блока
    assert "```" not in result["report"]
    assert '"verdict"' not in result["report"]
    assert "Инструкция проверена" in result["report"]
    # check ходил в llm с системным промтом, дополненным EDITS_INSTRUCTION
    assert len(calls) == 1
    assert calls[0]["system_prompt"].endswith(main.EDITS_INSTRUCTION)
    assert "охрана объектов" in calls[0]["user_message"]

    md = client.get(f"/api/jobs/{job_id}/export", params={"format": "md"})
    assert md.status_code == 200
    assert "Сводный отчёт DI_Check" in md.text
    assert "## di.docx" in md.text
    assert "Инструкция проверена" in md.text

    html = client.get(f"/api/jobs/{job_id}/export", params={"format": "html"})
    assert html.status_code == 200
    assert html.text.startswith("<!DOCTYPE html>")
    assert "Инструкция проверена" in html.text

    docx_resp = client.get(f"/api/jobs/{job_id}/export", params={"format": "docx"})
    assert docx_resp.status_code == 200
    assert docx_resp.headers["content-type"].startswith(DOCX_MIME)
    report = Document(io.BytesIO(docx_resp.content))
    assert "Сводный отчёт DI_Check" in "\n".join(p.text for p in report.paragraphs)
    row = report.tables[0].rows[1].cells
    assert [cell.text for cell in row] == ["di.docx", "Есть риски", "1", "done"]


# ---------- 3. Детерминированный /fix без LLM ----------

def test_fix_deterministic_applies_edits_without_llm(
    client, with_api_key, make_fake_llm, make_docx_bytes
):
    calls = make_fake_llm([structured_response()])
    resp = post_check(client, make_docx_bytes(BODY_PARAGRAPHS))
    assert resp.status_code == 200
    job_id = resp.json()["jobId"]
    poll_job(client, job_id, "done")
    assert len(calls) == 1  # один вызов от check

    fix = client.post(
        f"/api/jobs/{job_id}/fix", json={"resultIndex": 0, "editIds": ["e1"]}
    )
    assert fix.status_code == 200
    assert fix.headers["content-type"].startswith(DOCX_MIME)
    assert len(calls) == 1  # /fix не обращался к LLM

    text = docx_text(fix.content)
    assert REPLACEMENT in text
    assert QUOTE not in text


# ---------- 4. LLM-фолбэк /fix ----------

def test_fix_llm_fallback_when_quote_missing(
    client, with_api_key, make_fake_llm, make_docx_bytes
):
    calls = make_fake_llm(
        [
            structured_response(quote=MISSING_QUOTE),
            "## Исправленный документ\n\nНовый текст документа",
        ]
    )
    resp = post_check(client, make_docx_bytes(BODY_PARAGRAPHS))
    job_id = resp.json()["jobId"]
    data = poll_job(client, job_id, "done")
    assert [e["id"] for e in data["results"][0]["edits"]] == ["e1"]

    fix = client.post(
        f"/api/jobs/{job_id}/fix", json={"resultIndex": 0, "editIds": ["e1"]}
    )
    assert fix.status_code == 200
    text = docx_text(fix.content)
    assert "Исправленный документ" in text
    assert "Новый текст документа" in text
    assert len(calls) == 2  # check + фолбэк
    assert calls[1]["system_prompt"] == main.FIX_SYSTEM_PROMPT


# ---------- 5. Ошибки /fix ----------

def test_fix_validation_errors(client, make_job):
    make_job(
        job_id="jobfix001",
        edits=[
            {
                "id": "e1",
                "title": "t",
                "original": "цитата",
                "replacement": "замена",
                "reason": "r",
            }
        ],
        texts={0: "Текст с цитатой."},
    )
    empty = client.post("/api/jobs/jobfix001/fix", json={"resultIndex": 0, "editIds": []})
    assert empty.status_code == 400
    assert empty.json()["detail"] == "empty_edit_ids"

    unknown = client.post(
        "/api/jobs/jobfix001/fix", json={"resultIndex": 0, "editIds": ["nope"]}
    )
    assert unknown.status_code == 400
    assert "unknown_edit_ids" in unknown.json()["detail"]

    missing = client.post(
        "/api/jobs/missing000/fix", json={"resultIndex": 0, "editIds": ["e1"]}
    )
    assert missing.status_code == 404
    assert missing.json()["detail"] == "job_not_found"


# ---------- 6. Отмена ----------

def test_cancel_running_job_and_idempotent_repeat(
    client, with_api_key, monkeypatch, make_docx_bytes
):
    async def slow_chat_completion(provider, model, system_prompt, user_message):
        await asyncio.sleep(1.0)
        return "Отчёт."

    monkeypatch.setattr(main.llm, "chat_completion", slow_chat_completion)
    resp = post_check(client, make_docx_bytes(BODY_PARAGRAPHS))
    assert resp.status_code == 200
    job_id = resp.json()["jobId"]

    # ждём, пока задача файла реально начнёт выполняться (иначе гонка с cancel)
    deadline = time.monotonic() + 5.0
    while time.monotonic() < deadline:
        state = client.get(f"/api/jobs/{job_id}").json()
        if state["results"][0]["status"] == "running":
            break
        time.sleep(0.02)
    else:
        pytest.fail("задача проверки так и не начала выполняться")

    cancel = client.post(f"/api/jobs/{job_id}/cancel")
    assert cancel.status_code == 200
    assert cancel.json() == {"ok": True}

    data = poll_job(client, job_id, "cancelled")
    assert data["results"][0]["status"] == "cancelled"

    # повторная отмена завершённой джобы — идемпотентна
    again = client.post(f"/api/jobs/{job_id}/cancel")
    assert again.status_code == 200
    assert again.json() == {"ok": True, "status": "cancelled"}

    assert client.post("/api/jobs/nosuchjob/cancel").status_code == 404


# ---------- 7. Лимиты входа ----------

def test_check_limits(
    client, monkeypatch, with_api_key, make_fake_llm, make_docx_bytes
):
    make_fake_llm(["Отчёт."])

    monkeypatch.setattr(main, "MAX_FILES", 1)
    two_files = [
        ("files", ("a.docx", make_docx_bytes(["Текст А"]), DOCX_MIME)),
        ("files", ("b.docx", make_docx_bytes(["Текст Б"]), DOCX_MIME)),
    ]
    resp = client.post("/api/check", files=two_files)
    assert resp.status_code == 400
    assert "too_many_files" in resp.json()["detail"]

    monkeypatch.setattr(main, "MAX_FILE_SIZE", 10)
    resp = client.post(
        "/api/check", files={"files": ("big.txt", b"x" * 20, "text/plain")}
    )
    assert resp.status_code == 400
    assert "file_too_large" in resp.json()["detail"]

    resp = client.post(
        "/api/check",
        files={"files": ("prog.exe", b"MZ0000", "application/octet-stream")},
    )
    assert resp.status_code == 400
    assert "unsupported_format" in resp.json()["detail"]

    resp = client.post(
        "/api/check", files={"files": ("empty.txt", b"", "text/plain")}
    )
    assert resp.status_code == 400
    assert "empty_file" in resp.json()["detail"]


# ---------- 8. Magic-байты ----------

def test_check_rejects_fake_binary_format(client, with_api_key, make_fake_llm):
    make_fake_llm(["Отчёт."])
    resp = client.post(
        "/api/check",
        files={"files": ("x.docx", "это точно не docx, просто текст".encode(), DOCX_MIME)},
    )
    assert resp.status_code == 400
    assert resp.json()["detail"] == "invalid_file_format: x.docx"


def test_check_text_formats_skip_magic_check(
    client, with_api_key, make_fake_llm
):
    """Для .txt/.md проверки magic-байтов нет."""
    make_fake_llm(["Отчёт."])
    resp = client.post(
        "/api/check",
        files={"files": ("note.md", "совсем не magic, просто текст".encode(), "text/markdown")},
    )
    assert resp.status_code == 200
    poll_job(client, resp.json()["jobId"], "done")


# ---------- 9. Выгрузка просроченных джоб ----------

def test_expired_jobs_evicted_on_check(
    client, with_api_key, make_fake_llm, make_docx_bytes, make_job
):
    make_fake_llm(["Отчёт."])
    old_done = make_job(
        job_id="olddone01", status="done", created_at=time.monotonic() - 25 * 3600
    )
    old_running = make_job(
        job_id="oldrun001", status="running", created_at=time.monotonic() - 25 * 3600
    )

    resp = post_check(client, make_docx_bytes(BODY_PARAGRAPHS))
    assert resp.status_code == 200
    new_id = resp.json()["jobId"]

    assert old_done["id"] not in main._jobs
    assert old_running["id"] in main._jobs  # активные не трогаем
    assert new_id in main._jobs


# ---------- 10. GET /api/jobs/{id}: 404 и отсутствие служебных полей ----------

def test_get_job_missing_and_no_internal_fields(client, make_job):
    assert client.get("/api/jobs/missing00").status_code == 404

    job = make_job(job_id="svcjob001", texts={0: "СЕКРЕТНЫЙ ТЕКСТ"}, files={0: b"SECRETS"})
    job["tasks"] = ["task"]

    resp = client.get(f"/api/jobs/{job['id']}")
    assert resp.status_code == 200
    body = resp.json()
    assert set(body.keys()) == {"id", "status", "results"}
    raw = resp.text
    for secret in (
        "_texts",
        "_files",
        "tasks",
        "created_at",
        "cancel_requested",
        "СЕКРЕТНЫЙ ТЕКСТ",
        "SECRETS",
    ):
        assert secret not in raw


# ---------- 11. Обрезка длинного текста ----------

def test_long_text_truncated(
    client, with_api_key, make_fake_llm, make_docx_bytes, monkeypatch
):
    calls = make_fake_llm(["Отчёт."])
    monkeypatch.setattr(main, "MAX_TEXT_CHARS", 50)
    long_text = " ".join(f"слово{index}оченьдлинное" for index in range(20))
    assert len(long_text) > 50

    resp = post_check(client, make_docx_bytes([long_text]))
    assert resp.status_code == 200
    data = poll_job(client, resp.json()["jobId"], "done")

    assert data["results"][0]["textTruncated"] is True
    assert long_text[:50] in calls[0]["user_message"]
    assert long_text[50:] not in calls[0]["user_message"]
