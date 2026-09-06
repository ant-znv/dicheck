"""Сквозная проверка /api/check -> export?format=docx -> /fix-all против сервера на 8788.

Запуск: .venv/Scripts/python.exe -X utf8 backend/testdata/verify_batch.py
Артефакты: backend/testdata/batch_report.docx, backend/testdata/fixed_all.zip
"""
import io
import json
import sys
import time
import zipfile
from pathlib import Path

import httpx
from docx import Document

BASE = "http://127.0.0.1:8788"
TD = Path(__file__).resolve().parent
FILES = ["sample_di.docx", "sample_di_buhgalter.docx", "sample_di_menedzher.docx"]


def main() -> int:
    client = httpx.Client(base_url=BASE, timeout=60)

    # 1. /api/check с тремя файлами
    multipart = [
        ("files", (name, (TD / name).read_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        for name in FILES
    ]
    data = {
        "contractSubject": "Разработка ПО по госконтракту № 177-ГК",
        "employmentType": "full",
        "extraContext": "Требуется знание 44-ФЗ",
    }
    r = client.post("/api/check", files=multipart, data=data)
    assert r.status_code == 200, r.text
    job_id = r.json()["jobId"]
    print(f"jobId={job_id}")

    # 2. ждём done
    deadline = time.time() + 600
    job = None
    while time.time() < deadline:
        r = client.get(f"/api/jobs/{job_id}")
        job = r.json()
        st = [res["status"] for res in job["results"]]
        print(f"  statuses: {st}")
        if job["status"] == "done":
            break
        time.sleep(5)
    assert job and job["status"] == "done", "job not finished in time"
    for res in job["results"]:
        print(f"  {res['filename']}: {res['status']}, edits={len(res.get('edits') or [])}")
    assert all(res["status"] == "done" for res in job["results"]), "not all done"

    # 3. export?format=docx
    r = client.get(f"/api/jobs/{job_id}/export", params={"format": "docx"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ), r.headers["content-type"]
    report_path = TD / "batch_report.docx"
    report_path.write_bytes(r.content)
    print(f"batch_report.docx: {len(r.content)} bytes, CD={r.headers.get('content-disposition')}")

    # проверка структуры docx
    doc = Document(io.BytesIO(r.content))
    assert doc.tables, "no summary table"
    hdr = [c.text for c in doc.tables[0].rows[0].cells]
    print(f"  summary table header: {hdr}, rows={len(doc.tables[0].rows)}")
    assert len(doc.tables[0].rows) == 1 + len(FILES), "wrong summary rows"
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for name in FILES:
        assert any(name in h for h in heads), f"no section for {name}"
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Было" in full_text and "Будет" in full_text, "no edits section"
    print("  docx structure OK: title, table, per-file sections, edits with Было/Будет")

    # md/html не сломаны
    for fmt in ("md", "html"):
        r = client.get(f"/api/jobs/{job_id}/export", params={"format": fmt})
        assert r.status_code == 200 and len(r.content) > 100, fmt
    print("  md/html export OK")

    # 4. /fix-all с выбранными правками для 2 файлов
    items = []
    for idx in (0, 1):
        edits = job["results"][idx].get("edits") or []
        ids = [e["id"] for e in edits[:2]] or []
        assert ids, f"no edits for result {idx}"
        items.append({"resultIndex": idx, "editIds": ids})
    print(f"fix-all items: {json.dumps(items, ensure_ascii=False)}")
    r = client.post(f"/api/jobs/{job_id}/fix-all", json={"items": items}, timeout=600)
    assert r.status_code == 200, f"{r.status_code} {r.text}"
    assert r.headers["content-type"] == "application/zip", r.headers["content-type"]
    print(f"fix-all zip: {len(r.content)} bytes, CD={r.headers.get('content-disposition')}")
    zip_path = TD / "fixed_all.zip"
    zip_path.write_bytes(r.content)

    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = zf.namelist()
    print(f"  zip entries: {names}")
    docx_names = [n for n in names if n.endswith(".docx")]
    assert len(docx_names) == 2, docx_names
    for n in docx_names:
        d = Document(io.BytesIO(zf.read(n)))
        text = "\n".join(p.text for p in d.paragraphs)
        assert len(text) > 500, f"{n} looks empty"
        # хотя бы одна replacement-формулировка должна попасть в документ
        idx = 0 if "sample_di_исправленная" in n else 1
        applied = any(
            e["replacement"] and e["replacement"][:40] in text
            for e in job["results"][idx]["edits"]
            if e["id"] in items[idx]["editIds"]
        )
        print(f"  {n}: paragraphs={len(d.paragraphs)}, replacement applied={applied}")
        assert applied, f"{n}: replacement text not found"
    if "_errors.txt" in names:
        print(f"  _errors.txt: {zf.read('_errors.txt').decode('utf-8')}")

    # 5. ошибка 400 на пустом items
    r = client.post(f"/api/jobs/{job_id}/fix-all", json={"items": []})
    assert r.status_code == 400, f"expected 400, got {r.status_code}"
    print("empty items -> 400 OK")

    # resultIndex вне диапазона -> 400
    r = client.post(
        f"/api/jobs/{job_id}/fix-all",
        json={"items": [{"resultIndex": 99, "editIds": ["e1"]}]},
    )
    assert r.status_code == 400, f"expected 400, got {r.status_code}"
    print("resultIndex out of range -> 400 OK")

    # несуществующий jobId -> 404
    r = client.post("/api/jobs/deadbeef0000/fix-all", json={"items": items})
    assert r.status_code == 404, f"expected 404, got {r.status_code}"
    print("unknown job -> 404 OK")

    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
