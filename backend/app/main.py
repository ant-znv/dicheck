"""FastAPI-приложение DI_Check. Порт 8787, host 127.0.0.1."""
from __future__ import annotations

import asyncio
import copy
import html as html_mod
import io
import json
import logging
import re
import sys
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, urlsplit

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import extractors, history, llm, settings, updater
from .logsetup import setup_logging

# Логи настраивает точка входа (run.py); в dev (uvicorn напрямую) — консоль.
if not logging.getLogger().handlers:
    setup_logging(console=True)
logger = logging.getLogger("di_check")

app = FastAPI(title="DI_Check")

# CORS для dev (Vite dev server на 5173)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------- CSRF/DNS-rebinding защита ----------
# Сервер слушает только 127.0.0.1, поэтому и Origin, и Host должны указывать
# на loopback: чужой Origin — CSRF («простые» POST без preflight летят мимо
# CORS), чужой Host без Origin — DNS-rebinding. Middleware добавлен ПОСЛЕ
# CORSMiddleware, чтобы быть внешним и отдавать 403 до обработки CORS.

_LOOPBACK_HOSTS = {"127.0.0.1", "localhost", "::1"}


def _hostname_of(value: str) -> str | None:
    """Hostname без схемы и порта: для Origin ('http://h:p') и Host ('h:p')."""
    if not value:
        return None
    raw = value if "//" in value else f"//{value}"
    try:
        return urlsplit(raw).hostname
    except ValueError:
        return None


@app.middleware("http")
async def _loopback_only(request: Request, call_next):
    origin = request.headers.get("origin")
    if origin is not None and _hostname_of(origin) not in _LOOPBACK_HOSTS:
        return JSONResponse(status_code=403, content={"detail": "forbidden_origin"})
    if _hostname_of(request.headers.get("host", "")) not in _LOOPBACK_HOSTS:
        return JSONResponse(status_code=403, content={"detail": "forbidden_host"})
    return await call_next(request)

# Джобы в памяти
_jobs: dict[str, dict] = {}
_semaphore = asyncio.Semaphore(3)
# Сильные ссылки на фоновые задачи (иначе create_task может быть собран GC)
_run_tasks: set[asyncio.Task] = set()

# Лимиты входа и хранения
MAX_FILES = 20
MAX_FILE_SIZE = 20 * 1024 * 1024  # байт на файл
MAX_TEXT_CHARS = 120_000          # предел текста ДИ в user-message
JOB_TTL_SECONDS = 24 * 3600       # сколько жить завершённой джобе
MAX_JOBS = 50                     # максимум джоб в памяти
FILE_TTL_SECONDS = 3600           # сколько хранить байты/текст завершённой джобы
MAX_TOTAL_BYTES = 500 * 1024 * 1024  # суммарный лимит payload-байтов по всем джобам

# Магические байты бинарных форматов (расширение должно соответствовать содержимому)
_FILE_MAGIC = {
    ".docx": b"PK\x03\x04",                      # zip-контейнер (как и .odt)
    ".odt": b"PK\x03\x04",
    ".doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",  # OLE2
    ".pdf": b"%PDF",
}

if getattr(sys, "frozen", False):
    # PyInstaller: ресурсы распакованы в sys._MEIPASS (в onedir — каталог _internal)
    FRONTEND_DIST = Path(sys._MEIPASS) / "frontend" / "dist"
else:
    FRONTEND_DIST = Path(__file__).resolve().parent.parent.parent / "frontend" / "dist"


# ---------- Модели запросов ----------

class SettingsUpdate(BaseModel):
    activeProvider: str | None = None
    activeModel: str | None = None
    systemPrompt: str | None = None
    updateRepo: str | None = None


class ApiKeyUpdate(BaseModel):
    provider: str
    apiKey: str


class UpdateTokenUpdate(BaseModel):
    token: str


class TestRequest(BaseModel):
    provider: str
    model: str | None = None


class FixRequest(BaseModel):
    resultIndex: int
    editIds: list[str]


class DocumentPatch(BaseModel):
    position: str | None = None
    department: str | None = None
    title: str | None = None
    notes: str | None = None


class ExtractMetaBatchRequest(BaseModel):
    documentIds: list[int] | None = None


class OverlapsRequest(BaseModel):
    documentIds: list[int] | None = None


# ---------- Settings ----------

@app.get("/api/settings")
async def get_settings():
    return settings.build_settings_response()


@app.put("/api/settings")
async def put_settings(body: SettingsUpdate):
    try:
        settings.update_settings(
            active_provider=body.activeProvider,
            active_model=body.activeModel,
            system_prompt=body.systemPrompt,
        )
        if body.updateRepo is not None:
            settings.save_update_repo(body.updateRepo)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return settings.build_settings_response()


@app.put("/api/settings/update-token")
async def put_update_token(body: UpdateTokenUpdate):
    """Токен GitHub (только для приватного репозитория). Пустая строка — удалить."""
    settings.save_update_token(body.token)
    return {"ok": True, "hasToken": bool(settings.get_update_config()["token"])}


@app.get("/api/update/check")
async def update_check():
    """Проверка обновлений. Не падает никогда — ошибки внутри ответа."""
    cfg = settings.get_update_config()
    return updater.check(repo=cfg["repo"], token=cfg["token"])


@app.post("/api/update/install")
async def update_install():
    """Скачать установщик новой версии и запустить его; он сам перезапустит приложение."""
    cfg = settings.get_update_config()
    try:
        return updater.install(repo=cfg["repo"], token=cfg["token"])
    except updater.UpdateError as e:
        raise HTTPException(status_code=502, detail=str(e))


@app.put("/api/settings/apikey")
async def put_apikey(body: ApiKeyUpdate):
    try:
        settings.save_api_key(body.provider, body.apiKey)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"ok": True, "hasApiKey": settings.has_api_key(body.provider)}


@app.post("/api/settings/test")
async def test_connection(body: TestRequest):
    if body.provider not in settings.PROVIDERS:
        raise HTTPException(status_code=400, detail=f"unknown provider: {body.provider}")
    return await llm.test_connection(body.provider, body.model)


# ---------- Проверки ----------

def _build_user_message(
    filename: str,
    text: str,
    contract_subject: str,
    employment_type: str,
    extra_context: str,
) -> str:
    parts = [f"Должностная инструкция (файл: {filename}):\n\n{text}"]
    extras = []
    if contract_subject:
        extras.append(f"Предмет госконтракта: {contract_subject}")
    if employment_type == "full":
        extras.append("Занятость по контракту: сотрудник работает полностью по контракту")
    elif employment_type == "partial":
        extras.append("Занятость по контракту: сотрудник работает по контракту частично")
    if extra_context:
        extras.append(f"Дополнительный контекст (требования к персоналу и т.п.): {extra_context}")
    if extras:
        parts.append("---\n" + "\n".join(extras))
    return "\n\n".join(parts)


def _summary_verdict(report: str) -> str:
    """Эвристический вердикт по тексту отчёта (fallback, если модель не отдала свой)."""
    low = report.lower()
    if "не соответствует" in low:
        return "fail"
    if "риск" in low:
        return "risk"
    return "ok"


# ---------- Структурированные правки (edits) ----------

# Фиксированная инструкция, дописываемая к системному промту при проверке.
# НЕ редактируется пользователем и не входит в default_prompt.md.
EDITS_INSTRUCTION = """

---

В КОНЦЕ своего ответа обязательно выведи fenced-блок ```json со сводным вердиктом и списком конкретных правок текста должностной инструкции в формате:

```json
{"verdict": "risk", "edits": [{"id": "e1", "title": "краткое название правки", "original": "дословная цитата из текста ДИ, которую нужно изменить", "replacement": "готовая формулировка для вставки вместо цитаты", "reason": "почему нужна правка (норма/риск)"}]}
```

где verdict — общий вердикт по всей ДИ: "ok" (соответствует), "risk" (есть риски) или "fail" (не соответствует).

Требования к правкам:
- original — точная дословная цитата фрагмента текста инструкции; если подходящего фрагмента нет (например, раздел нужно добавить), опиши место вставки;
- replacement — полная готовая формулировка для вставки;
- правки должны быть атомарными (одна правка — одно место в тексте), id — e1, e2, e3...;
- блок ```json должен быть САМЫМ ПОСЛЕДНИМ в ответе, после него ничего не пиши.
"""

_FENCED_RE = re.compile(r"```[ \t]*(\w*)[ \t]*\r?\n(.*?)```", re.DOTALL)
_EDIT_KEYS = ("id", "title", "original", "replacement", "reason")
_VALID_VERDICTS = {"ok", "risk", "fail"}

# ---------- Починка markdown-таблиц в отчёте ----------
# Модель иногда выдаёт таблицу проверок, склеенную в одну строку/блок —
# рендерер показывает её сплошным текстом (см. docs/api-contract.md).

_DELIM_CELL_RE = re.compile(r":?-{3,}:?")


def _table_cells(line: str) -> list[str] | None:
    """Ячейки таблицной строки (обязательны leading/trailing '|'), иначе None."""
    s = line.strip()
    if not s.startswith("|") or not s.endswith("|") or s.count("|") < 2:
        return None
    return [c.strip() for c in s[1:-1].split("|")]


def _rebuild_table_block(lines: list[str]) -> list[str] | None:
    """Разбивает склеенные ряды таблицного блока на отдельные строки.

    Ячейки всех строк блока складываются в один поток, ищется серия
    разделительных ячеек '---' (она задаёт число колонок), шапка и данные
    нарезаются заново. Возвращает None, если блок нельзя разобрать однозначно
    (нет разделителя, шапка не равна числу колонок, данные не кратны) —
    тогда исходный текст оставляется как есть. Пустые ячейки/разделители
    рядов в потоке неразличимы, поэтому выбрасываются: таблицы с пустыми
    ячейками обычно не проходят проверку кратности и не трогаются.
    """
    stream: list[str] = []
    for line in lines:
        stream.extend(c for c in (_table_cells(line) or []) if c)

    run_start = run_end = None
    i = 0
    while i < len(stream):
        if _DELIM_CELL_RE.fullmatch(stream[i]):
            j = i
            while j < len(stream) and _DELIM_CELL_RE.fullmatch(stream[j]):
                j += 1
            if run_start is not None:  # вторая серия '---' — неоднозначно
                return None
            run_start, run_end = i, j
            i = j
        else:
            i += 1
    if run_start is None or run_start == 0:
        return None
    col = run_end - run_start
    header = stream[:run_start]
    data = stream[run_end:]
    if len(header) != col or not data or len(data) % col != 0:
        return None

    out = ["| " + " | ".join(header) + " |"]
    out.append("| " + " | ".join(stream[run_start:run_end]) + " |")
    for r in range(0, len(data), col):
        row = data[r : r + col]
        if not any(row):
            return None
        out.append("| " + " | ".join(row) + " |")
    return out


def _repair_markdown_tables(report: str) -> str:
    """Чинит таблицы, склеенные моделью в одну строку или подряд идущие строки.

    Корректно размеченные таблицы проходят через пересборку без изменений
    (идемпотентно); неоднозначные блоки не трогаются вовсе.
    """
    out: list[str] = []
    block: list[str] = []

    def flush() -> None:
        if not block:
            return
        rebuilt = _rebuild_table_block(block)
        out.extend(rebuilt if rebuilt is not None else block)
        block.clear()

    for line in report.split("\n"):
        if _table_cells(line) is not None:
            block.append(line)
        else:
            flush()
            out.append(line)
    flush()
    return "\n".join(out)


def _normalize_edits(data: object) -> list[dict] | None:
    """Валидирует распарсенный JSON и приводит edits к списку dict'ов по контракту."""
    if not isinstance(data, dict) or not isinstance(data.get("edits"), list):
        return None
    edits = []
    for i, item in enumerate(data["edits"], 1):
        if not isinstance(item, dict):
            continue
        edit = {k: str(item.get(k) or "") for k in _EDIT_KEYS}
        if not edit["id"]:
            edit["id"] = f"e{i}"
        edits.append(edit)
    return edits


def _matching_brace(text: str, start: int) -> int:
    """Индекс символа после закрывающей '}' для '{' в позиции start (с учётом строк). -1, если не нашлось."""
    depth = 0
    in_str = False
    escape = False
    for i in range(start, len(text)):
        ch = text[i]
        if in_str:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return i + 1
    return -1


def _parse_edits(report: str) -> tuple[str, list[dict], str | None]:
    """Извлекает блок {"verdict": ..., "edits": [...]} из ответа модели.

    Возвращает (очищенный отчёт, список правок, вердикт | None). Вердикт берётся
    из structured-блока; если его нет — вызывающий код использует эвристику.
    Если парсинг не удался — (исходный отчёт, [], None).
    """
    # 1) последний fenced-блок, содержащий {"edits": ...}
    for m in reversed(list(_FENCED_RE.finditer(report))):
        try:
            data = json.loads(m.group(2).strip())
        except json.JSONDecodeError:
            continue
        edits = _normalize_edits(data)
        if edits is not None:
            verdict = data.get("verdict")
            return (
                (report[: m.start()] + report[m.end() :]).strip(),
                edits,
                verdict if verdict in _VALID_VERDICTS else None,
            )
    # 2) brace-matching от последнего '{"edits"'
    idx = report.rfind('{"edits"')
    if idx != -1:
        end = _matching_brace(report, idx)
        if end != -1:
            try:
                data = json.loads(report[idx:end])
            except json.JSONDecodeError:
                data = None
            edits = _normalize_edits(data)
            if edits is not None:
                verdict = data.get("verdict") if isinstance(data, dict) else None
                cleaned = (report[:idx] + report[end:]).strip()
                # убрать осиротевшие fence-ограждения вокруг удалённого блока
                cleaned = re.sub(r"```[ \t]*(?:json)?[ \t]*$", "", cleaned).strip()
                cleaned = re.sub(r"^[ \t]*```", "", cleaned).strip()
                return cleaned, edits, verdict if verdict in _VALID_VERDICTS else None
    return report, [], None


async def _check_file(
    job: dict,
    index: int,
    result: dict,
    data: bytes,
    user_fields: dict,
    state: dict,
) -> None:
    started = time.monotonic()
    try:
        if job.get("cancel_requested"):
            # отмена пришла до старта задачи (узкое окно между create_check и _run_job)
            result["status"] = "cancelled"
            return
        async with _semaphore:
            result["status"] = "running"
            try:
                text = await asyncio.to_thread(
                    extractors.extract_text, result["filename"], data
                )
            except extractors.ExtractionError as e:
                result["status"] = "error"
                result["error"] = str(e)
                logger.warning(
                    "Джоба %s, файл %s: извлечение текста не удалось: %s",
                    job["id"], result["filename"], e,
                )
                return
            if len(text) > MAX_TEXT_CHARS:
                text = text[:MAX_TEXT_CHARS]
                result["textTruncated"] = True
                logger.warning(
                    "Джоба %s, файл %s: текст обрезан до %d символов (слишком длинный)",
                    job["id"], result["filename"], MAX_TEXT_CHARS,
                )
            # исходный текст и байты хранятся в джобе (не отдаются через GET):
            # текст нужен для /fix, байты — для детерминированных правок .docx
            job["_texts"][index] = text
            job["_files"][index] = data
            user_message = _build_user_message(result["filename"], text, **user_fields)
            raw_report = await llm.chat_completion(
                state["activeProvider"],
                state["activeModel"],
                state["systemPrompt"] + EDITS_INSTRUCTION,
                user_message,
            )
            report, edits, verdict = _parse_edits(raw_report)
            report = _repair_markdown_tables(report)
            result["report"] = report
            result["edits"] = edits
            result["summaryVerdict"] = verdict or _summary_verdict(report)
            result["status"] = "done"
            # история: не отвечает за проверку — любая ошибка БД только в лог
            hist = job.get("_history")
            if hist and hist.get("batch_id") is not None:
                try:
                    hist["version_ids"][index] = history.save_check_version(
                        hist["batch_id"],
                        result["filename"],
                        Path(result["filename"]).suffix.lower(),
                        job["_files"].get(index, b""),
                        job["_texts"].get(index, ""),
                        bool(result.get("textTruncated")),
                        result.get("report"),
                        result.get("summaryVerdict"),
                        result.get("edits") or [],
                    )
                except Exception as e:
                    logger.warning("History save failed: %s", e)
            logger.info(
                "Джоба %s, файл %s: готово за %.1f с, правок %d, вердикт %s",
                job["id"], result["filename"], time.monotonic() - started,
                len(edits), result["summaryVerdict"],
            )
    except asyncio.CancelledError:
        # отмену не пробрасываем: gather соберёт задачу, джоба получит статус cancelled
        result["status"] = "cancelled"
        logger.info("Джоба %s, файл %s: проверка отменена", job["id"], result["filename"])
    except llm.LLMError as e:
        result["status"] = "error"
        result["error"] = str(e)
        logger.warning(
            "Джоба %s, файл %s: ошибка LLM: %s", job["id"], result["filename"], e
        )
    except Exception:
        logger.exception(
            "Джоба %s, файл %s: неожиданная ошибка проверки", job["id"], result["filename"]
        )
        result["status"] = "error"
        result["error"] = "Внутренняя ошибка проверки (подробности в логе сервера)"


async def _run_job(job_id: str, files: list[tuple[str, bytes]], user_fields: dict) -> None:
    job = _jobs[job_id]
    state = settings.get_settings_state()
    tasks = [
        asyncio.create_task(_check_file(job, index, result, data, user_fields, state))
        for index, (result, (_, data)) in enumerate(zip(job["results"], files))
    ]
    job["tasks"] = tasks
    try:
        # исключения _check_file обрабатывает сам; gather — страховка
        outcomes = await asyncio.gather(*tasks, return_exceptions=True)
    finally:
        job["tasks"] = []
    for out in outcomes:
        if isinstance(out, BaseException) and not isinstance(out, asyncio.CancelledError):
            logger.error("Джоба %s: задача проверки упала исключением: %r", job_id, out)
    if job.get("cancel_requested"):
        job["status"] = "cancelled"
    elif all(r["status"] == "error" for r in job["results"]):
        job["status"] = "error"
    else:
        job["status"] = "done"
    job["_finished_at"] = time.monotonic()
    _evict_job_payloads()
    logger.info("Джоба %s завершена, статус: %s", job_id, job["status"])


def _evict_expired_jobs() -> None:
    """Выгружает завершённые джобы старше TTL и сверх MAX_JOBS (активные не трогаем)."""
    now = time.monotonic()
    expired = [
        j for j in _jobs.values()
        if j["status"] != "running" and now - j["created_at"] > JOB_TTL_SECONDS
    ]
    for j in expired:
        del _jobs[j["id"]]
    if expired:
        logger.info("Выгружено просроченных джоб: %d", len(expired))
    while len(_jobs) > MAX_JOBS:
        finished = sorted(
            (j for j in _jobs.values() if j["status"] != "running"),
            key=lambda j: j["created_at"],
        )
        if not finished:
            break
        oldest = finished[0]
        del _jobs[oldest["id"]]
        logger.info(
            "Выгружена старая джоба %s (превышен лимит %d джоб)", oldest["id"], MAX_JOBS
        )


def _payload_bytes(job: dict) -> int:
    """Суммарный размер payload'ов джобы: байты файлов + тексты (utf-8)."""
    return (
        sum(len(data) for data in job["_files"].values())
        + sum(len(text.encode("utf-8")) for text in job["_texts"].values())
    )


def _evict_job_payloads() -> None:
    """Выгружает payload'ы завершённых джоб (байты файлов и тексты).

    Отчёты/results остаются на месте: GET /api/jobs/{id} и экспорт работают.
    (а) у джоб, завершённых дольше FILE_TTL_SECONDS назад; (б) пока суммарный
    объём payload-байтов превышает MAX_TOTAL_BYTES — у самых старых завершённых
    (по created_at). Запущенные джобы не трогаются.
    """
    now = time.monotonic()
    candidates = sorted(
        (
            j for j in _jobs.values()
            if j["status"] != "running" and (j["_files"] or j["_texts"])
        ),
        key=lambda j: j["created_at"],
    )
    total = sum(_payload_bytes(j) for j in _jobs.values())
    evicted: list[dict] = []
    for j in candidates:
        finished_at = j.get("_finished_at")
        expired = finished_at is not None and now - finished_at > FILE_TTL_SECONDS
        if expired or total > MAX_TOTAL_BYTES:
            evicted.append(j)
            total -= _payload_bytes(j)
    for j in evicted:
        j["_files"].clear()
        j["_texts"].clear()
    if evicted:
        logger.info("Выгружены payload'ы (байты/тексты) джоб: %d", len(evicted))


@app.post("/api/check")
async def create_check(
    files: list[UploadFile] = File(...),
    contractSubject: str = Form(""),
    employmentType: str = Form(""),
    extraContext: str = Form(""),
):
    state = settings.get_settings_state()
    if not settings.has_api_key(state["activeProvider"]):
        raise HTTPException(status_code=409, detail="no_api_key")

    if not files:
        raise HTTPException(status_code=400, detail="no_files")
    if len(files) > MAX_FILES:
        raise HTTPException(
            status_code=400,
            detail=f"too_many_files: {len(files)} (максимум {MAX_FILES})",
        )

    payloads: list[tuple[str, bytes]] = []
    results = []
    for f in files:
        ext = Path(f.filename or "").suffix.lower()
        if ext not in extractors.SUPPORTED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"unsupported_format: {f.filename}",
            )
        chunks: list[bytes] = []
        total_size = 0
        while True:
            chunk = await f.read(1024 * 1024)  # 1 МБ — читаем чанками, не целиком
            if not chunk:
                break
            total_size += len(chunk)
            if total_size > MAX_FILE_SIZE:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"file_too_large: {f.filename} "
                        f"(максимум {MAX_FILE_SIZE // (1024 * 1024)} МБ)"
                    ),
                )
            chunks.append(chunk)
        data = b"".join(chunks)
        if not data:
            raise HTTPException(status_code=400, detail=f"empty_file: {f.filename}")
        magic = _FILE_MAGIC.get(ext)
        if magic is not None and not data.startswith(magic):
            raise HTTPException(
                status_code=400,
                detail=f"invalid_file_format: {f.filename}",
            )
        payloads.append((f.filename, data))
        results.append(
            {
                "filename": f.filename,
                "status": "pending",
                "error": None,
                "report": None,
                "summaryVerdict": None,
                "edits": None,
                "textTruncated": False,
            }
        )

    _evict_expired_jobs()
    _evict_job_payloads()

    job_id = uuid.uuid4().hex[:12]
    # история: регистрируем батч проверки (упавшая БД не мешает проверке)
    history_batch_id: int | None = None
    try:
        history_batch_id = history.ensure_batch(
            job_id,
            state["activeProvider"],
            state["activeModel"],
            contractSubject,
            employmentType,
            extraContext,
        )
    except Exception as e:
        logger.warning("History save failed: %s", e)
    job = {
        "id": job_id,
        "status": "running",
        "created_at": time.monotonic(),
        "results": results,
        "_texts": {},  # index результата -> извлечённый текст (наружу не отдаётся)
        "_files": {},  # index результата -> исходные байты файла (для правок .docx)
        "tasks": [],   # активные задачи файлов (для отмены)
        "cancel_requested": False,
    }
    if history_batch_id is not None:
        job["_history"] = {"batch_id": history_batch_id, "version_ids": {}}
    _jobs[job_id] = job
    user_fields = {
        "contract_subject": contractSubject,
        "employment_type": employmentType,
        "extra_context": extraContext,
    }
    task = asyncio.create_task(_run_job(job_id, payloads, user_fields))
    _run_tasks.add(task)
    task.add_done_callback(_run_tasks.discard)
    logger.info(
        "Джоба %s запущена: файлов %d, провайдер %s, модель %s",
        job_id, len(files), state["activeProvider"], state["activeModel"],
    )
    return {"jobId": job_id}


@app.get("/api/jobs/{job_id}")
async def get_job(job_id: str):
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job_not_found")
    # служебные поля ("_texts", "_files", "tasks" и т.п.) наружу не отдаём
    return {"id": job["id"], "status": job["status"], "results": job["results"]}


@app.post("/api/jobs/{job_id}/cancel")
async def cancel_job(job_id: str):
    """Отменяет проверку: файлы в работе прерываются, ожидающие — получают cancelled."""
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job_not_found")
    if job["status"] != "running":
        return {"ok": True, "status": job["status"]}
    job["cancel_requested"] = True
    for task in list(job.get("tasks") or []):
        task.cancel()
    logger.info("Джоба %s: запрошена отмена", job_id)
    return {"ok": True}


# ---------- Применение правок (/fix) ----------

FIX_SYSTEM_PROMPT = """Ты — редактор кадровых документов. Тебе дают исходный текст должностной инструкции и список правок.

Твоя задача — внести ВСЕ перечисленные правки и вернуть ПОЛНЫЙ исправленный текст документа целиком.

Требования к ответу:
- верни ТОЛЬКО текст исправленного документа, без каких-либо комментариев, пояснений и преамбул;
- не сокращай и не опускай разделы документа — он должен остаться полным;
- оформи структуру в markdown: заголовок документа — "# ...", заголовки разделов — "## ...", перечисления — маркированные списки через "- ";
- не обрамляй ответ в ``` и не добавляй ничего после текста документа.
"""


def _clean_md_inline(text: str) -> str:
    """Убирает markdown-мусор из строки (**жирный**, `код`)."""
    return text.replace("**", "").replace("__", "").replace("`", "").strip()


def _parse_table_block(
    lines: list[str], start: int
) -> tuple[list[list[str]], int] | None:
    """GFM-таблица с позиции start: [(строки ячеек), индекс после таблицы] или None.

    Таблица = строка-шапка, за которой сразу идёт строка-разделитель '---'.
    """
    if start + 1 >= len(lines):
        return None
    header = _table_cells(lines[start])
    if not header:
        return None
    delim = _table_cells(lines[start + 1])
    if not delim or not all(_DELIM_CELL_RE.fullmatch(c) for c in delim):
        return None
    rows = [header]
    i = start + 2
    while i < len(lines):
        cells = _table_cells(lines[i])
        if not cells:
            break
        rows.append(cells)
        i += 1
    return rows, i


def _markdown_into_docx(doc: Document, text: str) -> None:
    """Добавляет markdown-подобный текст модели в существующий Document."""
    lines = text.strip().splitlines()
    # снять обрамление ```markdown ... ```, если модель всё же обернула ответ
    if lines and lines[0].strip().startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]

    i = 0
    while i < len(lines):
        line = lines[i]
        table = _parse_table_block(lines, i)
        if table:
            rows, i = table
            ncols = len(rows[0])
            docx_table = doc.add_table(rows=0, cols=ncols)
            docx_table.style = "Table Grid"
            for r, row in enumerate(rows):
                cells = docx_table.add_row().cells
                for c in range(ncols):
                    cell = cells[c]
                    cell.text = _clean_md_inline(row[c] if c < len(row) else "")
                    if r == 0:  # шапка — полужирным
                        for par in cell.paragraphs:
                            for run in par.runs:
                                run.bold = True
            continue
        if not line.strip():
            i += 1
            continue
        line = line.rstrip()
        if line.startswith("### "):
            doc.add_heading(_clean_md_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_clean_md_inline(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_clean_md_inline(line[2:]), level=1)
        elif re.match(r"^\s*[-*•]\s+", line):
            item = re.sub(r"^\s*[-*•]\s+", "", line)
            doc.add_paragraph(_clean_md_inline(item), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            item = re.sub(r"^\s*\d+[.)]\s+", "", line)
            doc.add_paragraph(_clean_md_inline(item), style="List Number")
        else:
            doc.add_paragraph(_clean_md_inline(line))
        i += 1


def _markdown_to_docx(text: str) -> bytes:
    """Простая генерация .docx из markdown-подобного текста модели."""
    doc = Document()
    _markdown_into_docx(doc, text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FixError(Exception):
    """Ошибка применения правок к одному файлу (для /fix-all — в _errors.txt)."""


# ---------- Детерминированное применение правок к исходному .docx ----------

def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_docx_paragraphs(doc: Document):
    """Все абзацы документа: тело, таблицы и непустые колонтитулы."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            try:
                if hf.is_linked_to_previous:
                    continue
                yield from hf.paragraphs
                for table in hf.tables:
                    yield from _iter_table_paragraphs(table)
            except Exception:
                continue


def _replace_in_paragraph(par: Paragraph, original: str, replacement: str) -> bool:
    """Заменяет original на replacement в абзаце, работая с runs.

    Форматирование первого затронутого run'а распространяется на replacement.
    Возвращает False, если original не найден в runs абзаца.
    """
    runs = par.runs
    joined = "".join(r.text for r in runs)
    idx = joined.find(original)
    if idx == -1:
        return False
    end = idx + len(original)
    pos = 0
    placed = False
    for r in runs:
        r_start = pos
        r_end = pos + len(r.text)
        pos = r_end
        if r_end <= idx or r_start >= end:
            continue  # run вне диапазона замены
        prefix = r.text[: max(0, idx - r_start)]
        suffix = r.text[max(0, end - r_start):]
        r.text = (prefix + replacement + suffix) if not placed else (prefix + suffix)
        placed = True
    return True


def _apply_edits_to_docx(
    data: bytes, edits: list[dict]
) -> tuple[Document | None, list[dict]]:
    """Применяет правки к копии исходного .docx, сохраняя форматирование.

    Возвращает (Document | None, правки, которые не удалось применить на уровне
    runs — их добирает LLM-фолбэк).
    """
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        logger.warning("Детерминированные правки отменены: .docx не открылся: %s", e)
        return None, list(edits)
    failed: list[dict] = []
    for e in edits:
        applied = False
        if e["original"]:
            for par in _iter_docx_paragraphs(doc):
                joined = "".join(r.text for r in par.runs)
                if e["original"] in joined and _replace_in_paragraph(
                    par, e["original"], e["replacement"]
                ):
                    applied = True
                    break
        if not applied:
            failed.append(e)
    return doc, failed


def _build_fix_message(source_text: str, edits: list[dict]) -> str:
    edits_desc = "\n".join(
        f'{i}. Заменить «{e["original"]}» на «{e["replacement"]}»'
        + (f' (основание: {e["reason"]})' if e["reason"] else "")
        for i, e in enumerate(edits, 1)
    )
    return (
        f"Исходный текст должностной инструкции:\n\n{source_text}\n\n"
        f"---\n\nПравки, которые нужно внести:\n{edits_desc}"
    )


async def _perform_fix(job: dict, result_index: int, edit_ids: list[str]) -> tuple[bytes, str]:
    """Применяет правки к ДИ. Возвращает (docx-байты, способ: "docx"|"text"|"llm").

    Правки с дословной уникальной цитатой применяются без LLM: для исходного
    .docx — прямо в документе (сохраняя форматирование), для остальных
    источников — пересборкой из обновлённого текста. Недетерминированные правки
    (вставки, цитаты не найдены) добирает LLM-фолбэк.
    """
    result = job["results"][result_index]

    if not edit_ids:
        raise FixError("empty_edit_ids")

    by_id = {e["id"]: e for e in (result.get("edits") or [])}
    unknown = [eid for eid in edit_ids if eid not in by_id]
    if unknown:
        raise FixError(f"unknown_edit_ids: {', '.join(unknown)}")
    chosen = [by_id[eid] for eid in edit_ids]

    if result_index not in job["_files"]:
        # payload (байты/текст) выгружен из памяти (_evict_job_payloads) —
        # правки по этой джобе больше невозможны
        raise HTTPException(status_code=410, detail="payload_expired")
    source_text = job.get("_texts", {}).get(result_index)
    if not source_text:
        raise FixError("no_source_text")

    # 1) детерминированная часть: дословная цитата, уникальная в исходном тексте
    rest: list[dict] = []
    updated_text = source_text
    exact_applied: list[dict] = []
    for e in chosen:
        if e["original"] and source_text.count(e["original"]) == 1:
            if e["original"] in updated_text:
                updated_text = updated_text.replace(e["original"], e["replacement"], 1)
                exact_applied.append(e)
            else:
                rest.append(e)  # предыдущая замена задела текст этой правки
        else:
            rest.append(e)

    docx_data = job.get("_files", {}).get(result_index)
    if (
        exact_applied
        and docx_data
        and Path(result["filename"]).suffix.lower() == ".docx"
    ):
        doc, failed = await asyncio.to_thread(_apply_edits_to_docx, docx_data, exact_applied)
        rest = failed + rest
        if doc is not None and not failed:
            buf = io.BytesIO()
            await asyncio.to_thread(doc.save, buf)
            return buf.getvalue(), "docx"
        # если часть правок не легла на уровень runs — добираем через LLM ниже

    if exact_applied and not rest:
        # исходник не .docx (или без байтов), зато все правки применились к тексту
        return await asyncio.to_thread(_markdown_to_docx, updated_text), "text"

    # 2) LLM-фолбэк: вставки и цитаты, которые не удалось применить дословно
    state = settings.get_settings_state()
    fixed_text = await llm.chat_completion(
        state["activeProvider"],
        state["activeModel"],
        FIX_SYSTEM_PROMPT,
        _build_fix_message(updated_text, rest),
    )
    return await asyncio.to_thread(_markdown_to_docx, fixed_text), "llm"


@app.post("/api/jobs/{job_id}/fix")
async def fix_document(job_id: str, body: FixRequest):
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job_not_found")
    if body.resultIndex < 0 or body.resultIndex >= len(job["results"]):
        raise HTTPException(status_code=404, detail="result_not_found")

    try:
        docx_bytes, method = await _perform_fix(job, body.resultIndex, body.editIds)
    except FixError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except llm.LLMError as e:
        raise HTTPException(status_code=502, detail=str(e))

    logger.info(
        "Fix %s[%d]: применено методом %s", job_id, body.resultIndex, method
    )
    filename = f"{Path(job['results'][body.resultIndex]['filename']).stem}_исправленная.docx"
    # история: сохраняем исправленную версию (ошибка БД не ломает фикс)
    hist = job.get("_history") or {}
    parent = hist.get("version_ids", {}).get(body.resultIndex)
    if parent is not None:
        try:
            history.save_fixed_version(parent, docx_bytes, method, body.editIds, filename)
        except Exception as e:
            logger.warning("History save failed: %s", e)
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
        },
    )


# ---------- Пакетное исправление (/fix-all) ----------

class FixAllItem(BaseModel):
    resultIndex: int
    editIds: list[str]


class FixAllRequest(BaseModel):
    items: list[FixAllItem]


async def _fix_one(
    sem: asyncio.Semaphore, job: dict, item: FixAllItem
) -> tuple[str, bytes | None, str | None]:
    """Исправляет один файл. Возвращает (имя файла в архиве, docx | None, ошибка | None)."""
    result = job["results"][item.resultIndex]
    arcname = f"{Path(result['filename']).stem}_исправленная.docx"
    async with sem:
        try:
            docx_bytes, method = await _perform_fix(job, item.resultIndex, item.editIds)
        except (FixError, llm.LLMError) as e:
            logger.warning(
                "Fix-all %s[%d]: не удалось: %s", job["id"], item.resultIndex, e
            )
            return arcname, None, str(e)
    logger.info(
        "Fix-all %s[%d]: применено методом %s", job["id"], item.resultIndex, method
    )
    # история: сохраняем исправленную версию (ошибка БД не ломает фикс)
    hist = job.get("_history") or {}
    parent = hist.get("version_ids", {}).get(item.resultIndex)
    if parent is not None:
        try:
            history.save_fixed_version(parent, docx_bytes, method, item.editIds, arcname)
        except Exception as e:
            logger.warning("History save failed: %s", e)
    return arcname, docx_bytes, None


@app.post("/api/jobs/{job_id}/fix-all")
async def fix_all_documents(job_id: str, body: FixAllRequest):
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job_not_found")
    if not body.items:
        raise HTTPException(status_code=400, detail="empty_items")
    for item in body.items:
        if item.resultIndex < 0 or item.resultIndex >= len(job["results"]):
            raise HTTPException(
                status_code=400, detail=f"result_index_out_of_range: {item.resultIndex}"
            )

    sem = asyncio.Semaphore(3)
    outcomes = await asyncio.gather(*[_fix_one(sem, job, item) for item in body.items])

    successes: list[tuple[str, bytes]] = []
    failures: list[str] = []
    for result_index, (arcname, docx_bytes, error) in zip(
        (i.resultIndex for i in body.items), outcomes
    ):
        if docx_bytes is not None:
            successes.append((arcname, docx_bytes))
        else:
            failures.append(f"{job['results'][result_index]['filename']}: {error}")

    if not successes:
        raise HTTPException(status_code=502, detail="no_files_fixed")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used: set[str] = set()
        for arcname, docx_bytes in successes:
            # не допускаем коллизий имён при одинаковых исходных файлах
            name = arcname
            n = 2
            while name in used:
                name = f"{Path(arcname).stem}_{n}.docx"
                n += 1
            used.add(name)
            zf.writestr(name, docx_bytes)
        if failures:
            zf.writestr("_errors.txt", "\n".join(failures) + "\n")
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="ispravlennye_di.zip"'},
    )


# ---------- Экспорт ----------

def _export_markdown(job: dict) -> str:
    lines = [f"# Сводный отчёт DI_Check\n", f"Джоб: `{job['id']}`\n"]
    for r in job["results"]:
        lines.append(f"\n---\n\n## {r['filename']}\n")
        if r["status"] == "error":
            lines.append(f"**Ошибка:** {r['error']}\n")
        elif r["status"] == "cancelled":
            lines.append("_Проверка отменена._\n")
        elif r["report"]:
            lines.append(r["report"] + "\n")
        else:
            lines.append("_Нет результата._\n")
    return "\n".join(lines)


def _export_html(job: dict) -> str:
    md = _export_markdown(job)
    return (
        "<!DOCTYPE html>\n<html lang=\"ru\">\n<head>\n"
        "<meta charset=\"utf-8\">\n"
        f"<title>Отчёт DI_Check {html_mod.escape(job['id'])}</title>\n"
        "<style>body{font-family:Segoe UI,Arial,sans-serif;max-width:960px;"
        "margin:2em auto;padding:0 1em;}"
        "pre{white-space:pre-wrap;word-wrap:break-word;}</style>\n"
        "</head>\n<body>\n<pre>"
        + html_mod.escape(md)
        + "</pre>\n</body>\n</html>"
    )


_VERDICT_LABELS = {"ok": "OK", "risk": "Есть риски", "fail": "Не соответствует"}


def _export_docx(job: dict) -> bytes:
    """Единый .docx-отчёт: титульный блок, сводная таблица, секции по файлам."""
    doc = Document()
    doc.add_heading("Сводный отчёт DI_Check", level=0)
    doc.add_paragraph(f"Дата формирования: {datetime.now():%d.%m.%Y %H:%M}")
    doc.add_paragraph(f"Файлов проверено: {len(job['results'])}")

    doc.add_heading("Сводная таблица", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells, ("Файл", "Вердикт", "Правок", "Статус")
    ):
        cell.text = text
    for r in job["results"]:
        cells = table.add_row().cells
        cells[0].text = r["filename"]
        cells[1].text = _VERDICT_LABELS.get(r.get("summaryVerdict") or "", "—")
        cells[2].text = str(len(r.get("edits") or []))
        cells[3].text = r["status"]

    for r in job["results"]:
        if r["status"] != "done":
            continue
        doc.add_heading(f"Отчёт по файлу: {r['filename']}", level=1)
        if r.get("report"):
            _markdown_into_docx(doc, r["report"])
        else:
            doc.add_paragraph("Нет отчёта.")
        edits = r.get("edits") or []
        if edits:
            doc.add_heading("Правки", level=2)
            for e in edits:
                doc.add_paragraph(
                    f"{e['id']}. {e['title']}", style="List Number"
                )
                if e.get("reason"):
                    doc.add_paragraph(f"Основание: {e['reason']}")
                doc.add_paragraph(f"Было: {e['original']}")
                doc.add_paragraph(f"Будет: {e['replacement']}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/api/jobs/{job_id}/export")
async def export_job(job_id: str, format: str = "md"):
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job_not_found")
    if format == "md":
        content = _export_markdown(job).encode("utf-8")
        media_type = "text/markdown; charset=utf-8"
        ext = "md"
    elif format == "html":
        content = _export_html(job).encode("utf-8")
        media_type = "text/html; charset=utf-8"
        ext = "html"
    elif format == "docx":
        content = await asyncio.to_thread(_export_docx, job)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ext = "docx"
    else:
        raise HTTPException(status_code=400, detail="unsupported export format")
    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="di_check_{job_id}.{ext}"'
        },
    )


# ---------- История проверок (SQLite) ----------

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# media_type по расширению сохранённого файла (fix-версии всегда .docx)
_VERSION_MIME = {
    ".docx": _DOCX_MIME,
    ".doc": "application/msword",
    ".odt": "application/vnd.oasis.opendocument.text",
    ".pdf": "application/pdf",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/markdown; charset=utf-8",
}
META_SYSTEM_PROMPT = (
    "Извлеки из фрагмента должностной инструкции должность и подразделение. "
    'Верни только JSON без пояснений и markdown: {"position": "…", "department": "…"}. '
    "Если поле не найдено — пустая строка."
)
META_TEXT_LIMIT = 6000  # символов текста ДИ, уходящих в user-message extract-meta


def _validate_paging(limit: int, offset: int) -> None:
    if limit < 1 or limit > 500:
        raise HTTPException(status_code=400, detail="invalid_limit")
    if offset < 0:
        raise HTTPException(status_code=400, detail="invalid_offset")


def _document_summary(row: dict) -> dict:
    last = None
    if row.get("last_check_at") is not None:
        last = {
            "verdict": row.get("last_verdict"),
            "checkedAt": row["last_check_at"],
            "findingsCount": row.get("last_findings_count") or 0,
        }
    return {
        "id": row["id"],
        "title": row["title"],
        "position": row["position"],
        "department": row["department"],
        "notes": row["notes"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
        "versionsCount": row["versions_count"],
        "fixedCount": row["fixed_count"],
        "lastCheck": last,
    }


def _findings_count(v: dict) -> int:
    """Сводка несёт готовый count, детальная версию — список findings."""
    if "findings_count" in v:
        return int(v["findings_count"])
    return len(v.get("findings") or [])


def _version_summary(v: dict) -> dict:
    return {
        "id": v["id"],
        "origin": v["origin"],
        "filename": v["filename"],
        "createdAt": v["created_at"],
        "verdict": v["verdict"],
        "findingsCount": _findings_count(v),
        "fixMethod": v["fix_method"],
        "appliedEditIds": v["applied_edit_ids"],
        "parentVersionId": v["parent_version_id"],
        "jobId": v["job_id"],
    }


def _version_detail(v: dict) -> dict:
    return {
        **_version_summary(v),
        "documentId": v["document_id"],
        "text": v["text"],
        "textTruncated": v["text_truncated"],
        "report": v["report"],
        "findings": [
            {
                "editId": f["edit_id"],
                "title": f["title"],
                "original": f["original"],
                "replacement": f["replacement"],
                "reason": f["reason"],
            }
            for f in v["findings"]
        ],
    }


@app.get("/api/history/documents")
async def history_documents(search: str = "", limit: int = 100, offset: int = 0):
    _validate_paging(limit, offset)
    total, rows = history.list_documents(search, limit, offset)
    return {"total": total, "items": [_document_summary(r) for r in rows]}


@app.get("/api/history/documents/{document_id}")
async def history_document_get(document_id: int):
    doc = history.get_document(document_id)
    if doc is None:
        raise HTTPException(status_code=404, detail="document_not_found")
    return {
        "id": doc["id"],
        "title": doc["title"],
        "position": doc["position"],
        "department": doc["department"],
        "notes": doc["notes"],
        "createdAt": doc["created_at"],
        "updatedAt": doc["updated_at"],
        "versions": [_version_summary(v) for v in doc["versions"]],
    }


@app.patch("/api/history/documents/{document_id}")
async def history_document_patch(document_id: int, body: DocumentPatch):
    doc = history.update_document(document_id, **body.model_dump(exclude_none=True))
    if doc is None:
        raise HTTPException(status_code=404, detail="document_not_found")
    summary = history.document_summary(document_id)
    return _document_summary(summary)


@app.delete("/api/history/documents/{document_id}")
async def history_document_delete(document_id: int):
    if not history.delete_document(document_id):
        raise HTTPException(status_code=404, detail="document_not_found")
    return {"ok": True}


@app.get("/api/history/versions/{version_id}")
async def history_version_get(version_id: int):
    version = history.get_version(version_id)
    if version is None:
        raise HTTPException(status_code=404, detail="version_not_found")
    return _version_detail(version)


@app.get("/api/history/versions/{version_id}/download")
async def history_version_download(version_id: int):
    version = history.get_version(version_id)
    if version is None:
        raise HTTPException(status_code=404, detail="version_not_found")
    data = history.get_version_bytes(version_id)
    if not data:
        raise HTTPException(status_code=404, detail="no_file_data")
    ext = (version.get("file_ext") or "").lower()
    filename = version.get("filename") or "document.docx"
    return Response(
        content=data,
        media_type=_VERSION_MIME.get(ext, _DOCX_MIME),
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
        },
    )


def _parse_meta_response(raw: str) -> dict:
    """Разбирает ответ модели в {"position", "department"}; иначе LLMError."""
    data = None
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        idx = raw.find("{")
        if idx != -1:
            end = _matching_brace(raw, idx)
            if end != -1:
                try:
                    data = json.loads(raw[idx:end])
                except json.JSONDecodeError:
                    data = None
    if not isinstance(data, dict):
        raise llm.LLMError("invalid_meta_response")
    position = data.get("position")
    department = data.get("department")
    if not isinstance(position, str) or not isinstance(department, str):
        raise llm.LLMError("invalid_meta_response")
    return {"position": position.strip(), "department": department.strip()}


async def _extract_meta_for_document(document_id: int) -> tuple[str, str]:
    """LLM-извлечение должности/подразделения + запись в документ.

    HTTPException — нет документа/текста; llm.LLMError — сбой вызова/ответа.
    """
    if history.get_document(document_id) is None:
        raise HTTPException(status_code=404, detail="document_not_found")
    text = history.latest_check_text(document_id)
    if not text:
        raise HTTPException(status_code=409, detail="no_text")
    state = settings.get_settings_state()
    raw = await llm.chat_completion(
        state["activeProvider"],
        state["activeModel"],
        META_SYSTEM_PROMPT,
        text[:META_TEXT_LIMIT],
    )
    meta = _parse_meta_response(raw)
    updated = history.extract_meta_update(document_id, meta["position"], meta["department"])
    if updated is None:  # документ удалён параллельным запросом
        raise HTTPException(status_code=404, detail="document_not_found")
    return updated


@app.post("/api/history/documents/{document_id}/extract-meta")
async def history_extract_meta(document_id: int):
    try:
        position, department = await _extract_meta_for_document(document_id)
    except llm.LLMError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return {"position": position, "department": department}


@app.post("/api/history/extract-meta")
async def history_extract_meta_batch(body: ExtractMetaBatchRequest):
    document_ids = (
        body.documentIds if body.documentIds is not None
        else history.document_ids_missing_meta()
    )
    sem = asyncio.Semaphore(2)

    async def one(document_id: int) -> dict:
        try:
            async with sem:
                position, department = await _extract_meta_for_document(document_id)
            return {
                "documentId": document_id,
                "position": position,
                "department": department,
                "error": None,
            }
        except HTTPException as e:
            return {
                "documentId": document_id,
                "position": "",
                "department": "",
                "error": str(e.detail),
            }
        except Exception as e:
            return {
                "documentId": document_id,
                "position": "",
                "department": "",
                "error": str(e),
            }

    results = await asyncio.gather(*(one(i) for i in document_ids))
    return {"results": list(results)}


# ---------- Обязанности и пересечения (Ф1) ----------

DUTIES_SYSTEM_PROMPT = (
    "Извлеки из текста должностной инструкции список конкретных должностных обязанностей "
    '(разделы "Обязанности"/"Функции"/"Должностные обязанности"). '
    'Пропускай общие формулы ("соблюдает трудовую дисциплину", "выполняет приказы руководителя"). '
    'Верни только JSON без markdown: {"duties": ["обязанность 1", "обязанность 2"]}. '
    "Не больше 40 пунктов, каждый — одна фраза до 200 символов."
)
DUTIES_TEXT_LIMIT = 60_000  # символов текста ДИ, уходящих в user-message extract-duties
DUTIES_MAX_COUNT = 40       # cap на число обязанностей (валидация ответа модели)
DUTIES_MAX_LEN = 300        # cap на длину одной обязанности (валидация ответа модели)

OVERLAPS_SYSTEM_PROMPT = (
    "Ты находишь пересечения должностных обязанностей между сотрудниками РАЗНЫХ подразделений. "
    "Тебе даны документы: подразделение, должность, список обязанностей. "
    "Найди содержательные пересечения — одну и ту же фактическую работу, закреплённую за "
    "разными людьми/подразделениями (дублирование зон ответственности). "
    'НЕ включай общие формулировки и очевидные управленческие связи "подчинённый-руководитель" '
    "внутри одного подразделения. Пересечения ищи ТОЛЬКО между разными подразделениями. "
    "Верни только JSON без markdown: "
    '{"groups": [{"duty": "суть пересечения одной фразой", '
    '"comment": "кратко почему это проблема/что уточнить", '
    '"items": [{"documentId": 123, "duty": "формулировка из документа"}]}]}. '
    'Если пересечений нет — {"groups": []}. documentId указывай точно как дано.'
)
OVERLAP_CHUNK_SIZE = 8       # документов на один LLM-вызов сравнения
OVERLAP_LLM_CONCURRENCY = 2  # параллельных LLM-вызовов (извлечение и сравнение)


def _parse_json_object(raw: str) -> dict | None:
    """json.loads с fallback на brace-matching от первого '{'; иначе None."""
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        idx = raw.find("{")
        if idx == -1:
            return None
        end = _matching_brace(raw, idx)
        if end == -1:
            return None
        try:
            data = json.loads(raw[idx:end])
        except json.JSONDecodeError:
            return None
    return data if isinstance(data, dict) else None


def _parse_duties_response(raw: str) -> list[str]:
    """Разбирает ответ модели в {"duties": [...]}; иначе llm.LLMError."""
    data = _parse_json_object(raw)
    if data is None or not isinstance(data.get("duties"), list):
        raise llm.LLMError("invalid_duties_response")
    duties: list[str] = []
    for item in data["duties"]:
        if not isinstance(item, str):
            continue
        text = item.strip()
        if text:
            duties.append(text[:DUTIES_MAX_LEN])
    if not duties:
        raise llm.LLMError("invalid_duties_response")
    return duties[:DUTIES_MAX_COUNT]


async def _duties_via_llm(document_id: int, version_id: int, text: str) -> list[str]:
    """Один LLM-вызов извлечения обязанностей + сохранение в БД."""
    state = settings.get_settings_state()
    raw = await llm.chat_completion(
        state["activeProvider"],
        state["activeModel"],
        DUTIES_SYSTEM_PROMPT,
        text[:DUTIES_TEXT_LIMIT],
    )
    duties = _parse_duties_response(raw)
    history.set_duties(document_id, duties, version_id)
    return duties


async def _extract_duties_for_document(document_id: int) -> tuple[list[str], int]:
    """LLM-извлечение обязанностей + запись в документ (для endpoint'а).

    HTTPException — нет документа/текста; llm.LLMError — сбой вызова/ответа.
    """
    if history.get_document(document_id) is None:
        raise HTTPException(status_code=404, detail="document_not_found")
    check = history.latest_check_version(document_id)
    if check is None:
        raise HTTPException(status_code=409, detail="no_text")
    duties = await _duties_via_llm(document_id, *check)
    return duties, check[0]


@app.post("/api/history/documents/{document_id}/extract-duties")
async def history_extract_duties(document_id: int):
    try:
        duties, _version_id = await _extract_duties_for_document(document_id)
    except llm.LLMError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return {"documentId": document_id, "count": len(duties), "duties": duties}


def _parse_overlap_response(raw: str, chunk: list[dict]) -> list[dict]:
    """Разбирает ответ сравнения одного чанка.

    Оставляет только items с documentId из чанка; группы с менее чем двумя
    валидными items из РАЗНЫХ подразделений выбрасываются.
    """
    data = _parse_json_object(raw)
    if data is None or not isinstance(data.get("groups"), list):
        raise llm.LLMError("invalid_overlap_response")
    by_id = {e["documentId"]: e for e in chunk}
    groups: list[dict] = []
    for raw_group in data["groups"]:
        if not isinstance(raw_group, dict) or not isinstance(raw_group.get("items"), list):
            continue
        items = []
        for it in raw_group["items"]:
            if not isinstance(it, dict):
                continue
            doc_id = it.get("documentId")
            if not isinstance(doc_id, int) or doc_id not in by_id:
                continue  # documentId вне чанка / не число
            entry = by_id[doc_id]
            duty = it.get("duty")
            items.append({
                "documentId": doc_id,
                "position": entry["position"],
                "department": entry["department"],
                "duty": duty.strip() if isinstance(duty, str) else "",
            })
        departments = {(by_id[i["documentId"]]["department"] or "") for i in items}
        if len(items) < 2 or len(departments) < 2:
            continue
        groups.append({
            "duty": str(raw_group.get("duty") or "").strip(),
            "comment": str(raw_group.get("comment") or "").strip(),
            "items": items,
        })
    return groups


def _overlap_chunk_message(chunk: list[dict]) -> str:
    """user-message сравнения: перечисление документов с их обязанностями."""
    parts = []
    for e in chunk:
        lines = "\n".join(f"- {d}" for d in e["duties"][:DUTIES_MAX_COUNT])
        parts.append(
            f"### Документ {e['documentId']} — подразделение:"
            f" {e['department'] or '(не указано)'} — должность:"
            f" {e['position'] or '(не указано)'}\n{lines}"
        )
    return "\n\n".join(parts)


def _overlap_report_response(record: dict) -> dict:
    """Единая форма ответа overlaps (POST и GET): id + createdAt + результат."""
    return {"id": record["id"], "createdAt": record["created_at"], **record["result"]}


@app.post("/api/history/overlaps")
async def history_overlaps(body: OverlapsRequest):
    # documentIds null — все документы; иначе список (дубликаты схлопываем)
    if body.documentIds is None:
        ids = history.all_document_ids()
    else:
        ids = list(dict.fromkeys(body.documentIds))
    if not ids:
        raise HTTPException(status_code=400, detail="no_documents")

    docs: dict[int, dict] = {}
    for doc_id in ids:
        doc = history.get_document(doc_id)
        if doc is not None:
            docs[doc_id] = doc

    # 1. Гарантируем обязанности каждому документу: кэш, иначе LLM-извлечение
    sem = asyncio.Semaphore(OVERLAP_LLM_CONCURRENCY)

    async def ensure(duties_id: int) -> tuple[int, list[str] | None, str | None]:
        if duties_id not in docs:
            return duties_id, None, "document_not_found"
        duties, vid, _ts = history.get_duties(duties_id)
        check = history.latest_check_version(duties_id)
        if duties and check is not None and vid == check[0]:
            return duties_id, duties, None  # кэш актуален
        if check is None:
            return duties_id, None, "no_text"
        try:
            async with sem:
                duties = await _duties_via_llm(duties_id, *check)
            return duties_id, duties, None
        except llm.LLMError as e:
            return duties_id, None, f"извлечение не удалось: {e}"
        except Exception as e:
            return duties_id, None, f"извлечение не удалось: {e}"

    outcomes = await asyncio.gather(*(ensure(i) for i in ids))
    entries: list[dict] = []
    skipped: list[dict] = []
    for doc_id, duties, error in outcomes:
        doc = docs.get(doc_id) or {}
        base = {
            "documentId": doc_id,
            "position": doc.get("position") or "",
            "department": doc.get("department") or "",
        }
        if error:
            skipped.append({**base, "error": error})
        else:
            entries.append({**base, "duties": duties or []})

    # 2. Сравнение чанками по ≤8 документов: один LLM-вызов на чанк
    chunks = [
        entries[i : i + OVERLAP_CHUNK_SIZE]
        for i in range(0, len(entries), OVERLAP_CHUNK_SIZE)
    ]
    compare_sem = asyncio.Semaphore(OVERLAP_LLM_CONCURRENCY)

    async def compare(chunk: list[dict]) -> list[dict]:
        state = settings.get_settings_state()
        async with compare_sem:
            raw = await llm.chat_completion(
                state["activeProvider"],
                state["activeModel"],
                OVERLAPS_SYSTEM_PROMPT,
                _overlap_chunk_message(chunk),
            )
        return _parse_overlap_response(raw, chunk)

    groups: list[dict] = []
    if chunks:
        for part in await asyncio.gather(*(compare(c) for c in chunks)):
            groups.extend(part)

    result = {
        "documents": [
            {
                "documentId": e["documentId"],
                "position": e["position"],
                "department": e["department"],
                "dutiesCount": len(e["duties"]),
            }
            for e in entries
        ],
        "groups": groups,
        "skipped": skipped,
    }
    history.save_overlap_report([e["documentId"] for e in entries], result)
    return _overlap_report_response(history.latest_overlap_report())


@app.get("/api/history/overlaps")
async def history_overlaps_get():
    record = history.latest_overlap_report()
    if record is None:
        raise HTTPException(status_code=404, detail="no_report")
    return _overlap_report_response(record)


# ---------- Шаблон отчёта и сборка docx (Ф2, без LLM) ----------

TEMPLATE_KEYS = ("ДОЛЖНОСТЬ", "ПОДРАЗДЕЛЕНИЕ", "НАЗВАНИЕ", "ТЕКСТ", "ДАТА", "ЗАМЕТКИ")
_TEMPLATE_MAGIC = b"PK\x03\x04"


def _iter_template_paragraphs(doc: Document):
    """Абзацы шаблона: тело документа и таблицы (включая вложенные)."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)


def _template_keys(doc: Document) -> list[str]:
    """Ключи TEMPLATE_KEYS, встречающиеся в шаблоне ({{КЛЮЧ}})."""
    found = []
    for key in TEMPLATE_KEYS:
        marker = "{{" + key + "}}"
        for par in _iter_template_paragraphs(doc):
            if marker in "".join(r.text for r in par.runs):
                found.append(key)
                break
    return found


@app.put("/api/history/template")
async def history_template_put(request: Request):
    """Загрузка шаблона: тело запроса — сырые байты .docx (octet-stream)."""
    data = await request.body()
    if not data.startswith(_TEMPLATE_MAGIC):
        raise HTTPException(status_code=400, detail="invalid_template")
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise HTTPException(status_code=400, detail="invalid_template")
    history.save_template(data)
    return {"ok": True, "size": len(data), "keys": _template_keys(doc)}


@app.get("/api/history/template")
async def history_template_get():
    data = history.get_template()
    if data is None:
        return {"exists": False}
    try:
        keys = _template_keys(Document(io.BytesIO(data)))
    except Exception:
        keys = []
    return {"exists": True, "size": len(data), "keys": keys}


@app.delete("/api/history/template")
async def history_template_delete():
    history.delete_template()
    return {"ok": True}


def _set_paragraph_text(par: Paragraph, new_text: str) -> None:
    """Перезаписывает текст абзаца с сохранением форматирования.

    Текст кладётся в первый непустой run (его форматирование распространяется
    на всё значение), остальные run'ы очищаются.
    """
    target = None
    for r in par.runs:
        if target is None and r.text:
            target = r
        else:
            r.text = ""
    if target is None:
        target = par.runs[0] if par.runs else par.add_run("")
    target.text = new_text


def _insert_body_after_anchor(
    anchor: Paragraph, body_docx_bytes: bytes | None, text: str
) -> None:
    """Заменяет якорный абзац {{ТЕКСТ}} содержимым исправленного docx.

    Если body_docx_bytes передан и открывается — его тело (кроме sectPr)
    вставляется после якоря, сам якорь удаляется; иначе текст вставляется
    построчно абзацами со стилем якоря.
    """
    body_doc = None
    if body_docx_bytes:
        try:
            body_doc = Document(io.BytesIO(body_docx_bytes))
        except Exception as e:
            logger.warning("Шаблонная сборка: тело не открылось, вставляю текст: %s", e)
    if body_doc is not None:
        cursor = anchor._p
        for child in list(body_doc.element.body):
            if child.tag == qn("w:sectPr"):
                continue
            new_el = copy.deepcopy(child)
            cursor.addnext(new_el)
            cursor = new_el
    else:
        for line in (text or "").splitlines():
            anchor.insert_paragraph_before(line, style=anchor.style)
    anchor._p.getparent().remove(anchor._p)


def assemble_from_template(
    template_bytes: bytes,
    position: str = "",
    department: str = "",
    title: str = "",
    notes: str = "",
    body_docx_bytes: bytes | None = None,
    text: str = "",
) -> bytes:
    """Детерминированная сборка docx из шаблона с плейсхолдерами {{КЛЮЧ}} (без LLM).

    Простые ключи заменяются по всему тексту абзаца (тело + таблицы) с
    сохранением форматирования; {{ДАТА}} — текущая дата ДД.ММ.ГГГГ; абзац с
    {{ТЕКСТ}} заменяется содержимым body_docx_bytes либо построчным текстом.
    """
    doc = Document(io.BytesIO(template_bytes))
    values = {
        "ДОЛЖНОСТЬ": position,
        "ПОДРАЗДЕЛЕНИЕ": department,
        "НАЗВАНИЕ": title,
        "ЗАМЕТКИ": notes,
        "ДАТА": datetime.now().strftime("%d.%m.%Y"),
    }
    anchor = None
    for par in _iter_template_paragraphs(doc):
        joined = "".join(r.text for r in par.runs)
        if "{{" not in joined:
            continue
        if "{{ТЕКСТ}}" in joined:
            if anchor is None:
                anchor = par  # содержимое вставляется в первый якорь
            else:
                _set_paragraph_text(par, joined.replace("{{ТЕКСТ}}", ""))
            continue
        new_text = joined
        for key, value in values.items():
            new_text = new_text.replace("{{" + key + "}}", value)
        if new_text != joined:
            _set_paragraph_text(par, new_text)
    if anchor is not None:
        _insert_body_after_anchor(anchor, body_docx_bytes, text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_FILENAME_BAD_RE = re.compile(r'[\\/:*?"<>|\r\n]+')


def _safe_arc_stem(name: str) -> str:
    """Безопасное имя для архива: запрещённые в Windows символы → '_'."""
    cleaned = _FILENAME_BAD_RE.sub("_", (name or "")).strip(" .")
    return cleaned or "Документ"


async def _export_fixed_with_template(ids: list[int] | None) -> Response:
    """ZIP с исправленными ДИ, собранными из пользовательского шаблона."""
    template_bytes = history.get_template()
    if template_bytes is None:
        raise HTTPException(status_code=409, detail="no_template")
    details = history.latest_fixed_details(ids)
    if not details:
        raise HTTPException(status_code=404, detail="no_fixed_files")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used: set[str] = set()
        for d in details:
            data = await asyncio.to_thread(
                assemble_from_template,
                template_bytes,
                d["position"],
                d["department"],
                d["document_title"],
                d["notes"],
                d["file_bytes"],
                d["text"],
            )
            stem = _safe_arc_stem(d["position"] or d["document_title"])
            name = f"{stem}_исправленная.docx"
            n = 2
            while name in used:  # коллизии имён — как в /fix-all
                name = f"{stem}_исправленная_{n}.docx"
                n += 1
            used.add(name)
            zf.writestr(name, data)
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="ispravlennye_di.zip"'},
    )


@app.get("/api/history/export/fixed")
async def history_export_fixed(documentIds: str = "", template: bool = False):
    raw = (documentIds or "").strip()
    ids: list[int] | None
    if raw:
        try:
            ids = [int(part) for part in raw.split(",") if part.strip()]
        except ValueError:
            raise HTTPException(status_code=400, detail="invalid_document_ids")
    else:
        ids = None
    if template:
        return await _export_fixed_with_template(ids)
    files = history.latest_fixed_versions(ids)
    if not files:
        raise HTTPException(status_code=404, detail="no_fixed_files")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used: set[str] = set()
        for _title, data, filename in files:
            # не допускаем коллизий имён (как в /fix-all)
            name = filename or "document.docx"
            n = 2
            while name in used:
                name = f"{Path(name).stem}_{n}.docx"
                n += 1
            used.add(name)
            zf.writestr(name, data)
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="ispravlennye_di.zip"'},
    )


# ---------- Статика фронтенда (prod) ----------

if FRONTEND_DIST.is_dir():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIST), html=True), name="frontend")
