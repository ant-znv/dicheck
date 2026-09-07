"""Извлечение текста из файлов ДИ: .docx, .pdf, .txt/.md, .odt, .doc."""
from __future__ import annotations

import io
import logging
import re
import tempfile
from pathlib import Path

SUPPORTED_EXTENSIONS = {".docx", ".doc", ".pdf", ".txt", ".md", ".odt"}

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """Ошибка извлечения текста из файла."""


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        return _from_docx(data)
    if ext == ".pdf":
        return _from_pdf(data)
    if ext in (".txt", ".md"):
        return _from_text(data)
    if ext == ".odt":
        return _from_odt(data)
    if ext == ".doc":
        return _from_doc(data)
    raise ExtractionError(f"Неподдерживаемый формат файла: {ext or filename}")


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        header_lines, footer_lines = _docx_header_footer_lines(doc)
        parts = list(header_lines)
        parts.extend(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        parts.extend(footer_lines)
        text = "\n".join(parts).strip()
    except Exception as e:
        raise ExtractionError(f"Не удалось прочитать .docx: {e}")
    if not text:
        raise ExtractionError("Файл .docx не содержит текста")
    return text


def _docx_header_footer_lines(doc) -> tuple[list[str], list[str]]:
    """Строки колонтитулов всех секций; унаследованные (is_linked_to_previous) пропускаются."""
    headers: list[str] = []
    footers: list[str] = []
    for section in doc.sections:
        for target, part in (
            (headers, section.header),
            (headers, section.first_page_header),
            (headers, section.even_page_header),
            (footers, section.footer),
            (footers, section.first_page_footer),
            (footers, section.even_page_footer),
        ):
            try:
                if part is None or part.is_linked_to_previous:
                    continue
                target.extend(_docx_part_lines(part))
            except Exception:
                continue
    return _dedupe_lines(headers), _dedupe_lines(footers)


def _docx_part_lines(part) -> list[str]:
    """Строки части документа (тело или колонтитул): абзацы плюс строки таблиц."""
    lines = [p.text for p in part.paragraphs]
    for table in part.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return lines


def _dedupe_lines(lines: list[str]) -> list[str]:
    """Убрать повторные строки, сохраняя порядок первого вхождения."""
    seen: set[str] = set()
    unique: list[str] = []
    for line in lines:
        if line in seen:
            continue
        seen.add(line)
        unique.append(line)
    return unique


_PDF_OCR_HINT = (
    "Не удалось извлечь текст из .pdf (возможно, это скан — нужен OCR). "
    "Для распознавания сканов установите Tesseract OCR (языки rus+eng)"
)


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as e:
        raise ExtractionError(f"Не удалось прочитать .pdf: {e}")
    if text:
        return text
    return _from_pdf_ocr(data)


def _from_pdf_ocr(data: bytes) -> str:
    """OCR-фолбэк для .pdf: PyMuPDF рендерит страницы, Tesseract распознаёт (rus+eng)."""
    try:
        import fitz
        import pytesseract
        from PIL import Image
    except ImportError:
        raise ExtractionError(_PDF_OCR_HINT)
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        # Движок Tesseract не установлен или недоступен — распознавать нечем
        raise ExtractionError(_PDF_OCR_HINT)
    pages: list[str] = []
    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for page in pdf:
                try:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    pages.append(pytesseract.image_to_string(img, lang="rus+eng"))
                except Exception as e:
                    logger.warning("OCR страницы .pdf не удался: %s", e)
    except Exception as e:
        logger.warning("Не удалось открыть .pdf для OCR: %s", e)
    text = "\n".join(pages).strip()
    if not text:
        raise ExtractionError(
            "Не удалось извлечь текст из .pdf даже через OCR "
            "(возможно, файл повреждён или страницы пусты)"
        )
    return text


def _decode_text(data: bytes) -> str:
    """Декодировать байты в текст, перебирая типовые кодировки (BOM снимается автоматически)."""
    for encoding in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            text = data.decode(encoding).strip()
            if text:
                return text
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ExtractionError("Не удалось определить кодировку текстового файла")


def _from_text(data: bytes) -> str:
    return _decode_text(data)


def _from_odt(data: bytes) -> str:
    try:
        from odf.opendocument import load
        from odf import teletype, text as odftext

        doc = load(io.BytesIO(data))
        parts = []
        for elem in doc.getElementsByType(odftext.P) + doc.getElementsByType(
            odftext.H
        ):
            parts.append(teletype.extractText(elem))
        text = "\n".join(parts).strip()
    except Exception as e:
        raise ExtractionError(f"Не удалось прочитать .odt: {e}")
    if not text:
        raise ExtractionError("Файл .odt не содержит текста")
    return text


def _from_doc(data: bytes) -> str:
    """Старый бинарный .doc — через Word COM (pywin32)."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise ExtractionError(
            "pywin32 недоступен — сконвертируйте файл .doc в .docx и загрузите снова"
        )

    tmp_dir = Path(tempfile.mkdtemp(prefix="di_check_doc_"))
    src = tmp_dir / "input.doc"
    dst = tmp_dir / "output.txt"
    src.write_bytes(data)

    word = None
    try:
        pythoncom.CoInitialize()
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception as e:
            raise ExtractionError(
                "Не удалось запустить Microsoft Word для чтения .doc "
                f"({e}). Сконвертируйте файл в .docx и загрузите снова"
            )
        # Запрет макросов в недоверенных .doc: 3 = msoAutomationSecurityForceDisable
        try:
            word.AutomationSecurity = 3
        except Exception as e:
            logger.warning("Word AutomationSecurity недоступен (%s), продолжаем", e)
        try:
            word.DisplayAlerts = 0
        except Exception as e:
            logger.warning("Word DisplayAlerts недоступен (%s), продолжаем", e)
        word.Visible = False
        try:
            doc = word.Documents.Open(
                str(src), ReadOnly=True, ConfirmConversions=False,
                AddToRecentFiles=False, Visible=False,
            )
        except Exception as e:
            raise ExtractionError(
                f"Word не смог открыть .doc ({e}). Сконвертируйте файл в .docx"
            )
        try:
            doc.SaveAs2(str(dst), FileFormat=7)  # wdFormatUnicodeText
        finally:
            doc.Close(False)
        text = _decode_text(dst.read_bytes())
        if not text:
            raise ExtractionError("Файл .doc не содержит текста")
        return text
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
        for p in (src, dst):
            try:
                p.unlink()
            except OSError:
                pass
        try:
            tmp_dir.rmdir()
        except OSError:
            pass
